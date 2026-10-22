"""Checks of the master pipeline layer (src/pipeline/, src/ingest/pinnacle.py; docs/studies/master_pipeline/).

Fast and offline: no download, no build, no fit. They check what the notebook relies on:
    * every command the layer can issue is accepted by its target's own parser (dry runs, compared with --help);
    * the money arithmetic: odds conversion, proportional margin removal, side orientation, break-even roots;
    * the Pinnacle staging and the linker on the saved probe of 9 Oct 2026 (data/raw/pinnacle/probe/);
    * the runner: streaming, logs, receipts, failure, scrubbed environment;
    * the notebook: valid, no outputs, every cell compiles, every tp.<name> it calls exists.
"""
from __future__ import annotations

import datetime as dt
import json
import os
import re
import subprocess
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import src.pipeline as tp  # noqa: E402
from src.pipeline import bet_sheet as BS  # noqa: E402
from src.pipeline import config as CFG  # noqa: E402
from src.pipeline import diagnostics as D  # noqa: E402
from src.pipeline import live as L  # noqa: E402
from src.pipeline import runner as R  # noqa: E402

PROBE = ROOT / "data" / "raw" / "pinnacle" / "probe"
NOTEBOOK = ROOT / "scripts" / "master_pipeline.ipynb"
REPLAY = ROOT / "data" / "studies" / "daily_prediction" / "runs" / "replay_2026-07-13_R16_speed"
needs_probe = pytest.mark.skipif(not (PROBE / "tennis_matchups.json").exists(), reason="no Pinnacle probe")
needs_replay = pytest.mark.skipif(not (REPLAY / "predictions.csv").exists(), reason="no replay run")


@pytest.fixture()
def sandbox_out(tmp_path, monkeypatch):
    """Logs and receipts of the runner go to a temporary folder."""
    monkeypatch.setattr(CFG, "OUT", tmp_path)
    monkeypatch.setattr(CFG, "LOGS", tmp_path / "logs")
    monkeypatch.setattr(CFG, "RECEIPTS", tmp_path / "runs.jsonl")
    monkeypatch.setattr(R, "RUNNING", tmp_path / "running")
    return tmp_path


# ------------------------------------------------------------------------------------------------------------------
# interpreters and the runner

def test_interpreters_come_from_the_shims():
    for env, shim in (("tennis", "python.sh"), ("local", "python_tabpfn_local.sh"), ("api", "python_tabpfn.sh")):
        text = (ROOT / shim).read_text(encoding="utf-8")
        assert str(CFG.interpreter(env)).replace("\\", "/") in text.replace("\\", "/") or os.environ.get(
            {"tennis": "TENNIS_PYTHON", "local": "TENNIS_TABPFN_PYTHON", "api": "TENNIS_TABPFN_API_PYTHON"}[env])


def test_runner_streams_logs_and_receipts(sandbox_out, capsys):
    res = R.run_step(R.Step("echo", ["-c", "print('hello'); print('world')"]))
    assert res.ok and res.returncode == 0
    assert "hello" in capsys.readouterr().out
    assert "world" in Path(res.log).read_text(encoding="utf-8")
    rec = json.loads((sandbox_out / "runs.jsonl").read_text(encoding="utf-8").splitlines()[-1])
    assert rec["step"] == "echo" and rec["status"] == "ok" and rec["returncode"] == 0
    assert not list((sandbox_out / "running").glob("*.json"))          # the marker is gone once the step ended


def test_runner_raises_on_failure_with_the_tail(sandbox_out):
    with pytest.raises(R.StepFailed) as err:
        R.run_step(R.Step("fail", ["-c", "import sys; print('about to fail'); sys.exit(3)"]))
    assert err.value.result.returncode == 3 and "about to fail" in str(err.value)
    res = R.run_step(R.Step("fail_ok", ["-c", "import sys; sys.exit(1)"], ok_codes=(0, 1)))
    assert res.ok


def test_runner_dry_run_runs_nothing(sandbox_out, capsys):
    res = R.run_step(R.Step("would_fail", ["-c", "raise SystemExit(5)"]), dry_run=True)
    assert res.status == "dry-run" and res.returncode is None
    assert "DRY RUN" in capsys.readouterr().out


def test_child_environment_is_scrubbed(sandbox_out, monkeypatch, capsys):
    monkeypatch.setenv("DP_BASE_TREE", "C:/somewhere/else")
    monkeypatch.setenv("LLE_LAYER", "x.json")
    code = "import os; print('SEEN', os.environ.get('DP_BASE_TREE'), os.environ.get('LLE_LAYER'), os.environ.get('DP_WORK_TREE'))"
    R.run_step(R.Step("env", ["-c", code], extra_env={"DP_WORK_TREE": "C:/mine"}))
    out = capsys.readouterr().out
    assert "SEEN None None C:/mine" in out


def test_parallel_steps(sandbox_out):
    res = R.run_parallel([R.Step("a", ["-c", "print('a')"]), R.Step("b", ["-c", "print('b')"])])
    assert [r.status for r in res] == ["ok", "ok"]


# ------------------------------------------------------------------------------------------------------------------
# every command the layer issues exists in its target's parser

def _flags(help_text: str) -> set[str]:
    return set(re.findall(r"(--[a-z][a-z0-9-]*)", help_text))


def _help(args: list[str], env: str = "tennis") -> str:
    out = subprocess.run([str(CFG.interpreter(env)), *args, "--help"], cwd=ROOT, capture_output=True, text=True,
                         encoding="utf-8", errors="replace", env={**os.environ, "PYTHONIOENCODING": "utf-8"})
    return out.stdout + out.stderr


def _commands(monkeypatch) -> list[R.Step]:
    """Every step the notebook's functions would run, collected from their dry runs."""
    seen: list[R.Step] = []

    def capture(step, **kw):
        seen.append(step)
        return R.StepResult(step.name, step.shown(), "dry-run", None, 0.0, None)

    for mod in ("data", "train", "live", "predict", "thinking"):
        monkeypatch.setattr(f"src.pipeline.{mod}.run_step", capture, raising=False)
    monkeypatch.setattr("src.pipeline.runner.run_step", capture)
    monkeypatch.setattr("src.pipeline.train.run_parallel", lambda steps, **kw: [capture(s) for s in steps])
    monkeypatch.setattr("src.pipeline.train.run_steps", lambda steps, **kw: [capture(s) for s in steps])
    monkeypatch.setattr("src.pipeline.data.run_steps", lambda steps, **kw: [capture(s) for s in steps])
    monkeypatch.setattr("src.pipeline.runner.run_steps", lambda steps, **kw: [capture(s) for s in steps])
    monkeypatch.setattr("src.pipeline.data._refuse_paused", lambda *a, **k: None)
    tp.scheduled_downloads(only=["tml", "speeds"], dry_run=True)
    tp.accept_held_tml("2025_challenger.csv@78b86009de42", dry_run=True)
    tp.refresh_speeds(days_back=2, dry_run=True)
    tp.refresh_timeline_sources(dry_run=True)
    tp.refresh_odds_history(dry_run=True)
    tp.rebuild_shared_tables(confirm=True, dry_run=True)
    tp.build_instance("data/studies/current_best_model/x", sandbox="C:/tmp/x", dry_run=True)
    tp.finish_instance("data/studies/current_best_model/x", dry_run=True)
    tp.fit_seasons("template_2026-10-10", seasons="2026", dry_run=True)
    tp.compare_candidate("data/studies/current_best_model/x", dry_run=True)
    tp.check_history(dry_run=True)
    tp.make_checkpoint("2026-10-05", dry_run=True)
    tp.replay_check(dry_run=True)
    tp.grade(REPLAY, dry_run=True)
    tp.record_pinnacle(dry_run=True)
    tp.thinking.fit_season(tp.active_instance(), 2026, dry_run=True)
    return seen


def test_every_issued_command_parses(monkeypatch):
    steps = _commands(monkeypatch)
    assert len(steps) >= 25
    helps: dict[tuple, str] = {}
    for s in steps:
        args = [str(a) for a in s.args]
        if args[0] == "-m":
            target, rest = args[:2], args[2:]
        else:
            target, rest = args[:1], args[1:]
        sub = []
        if target[-1].endswith("dp.py") or target[-1] == "src.pipeline.thinking" or target[-1] == "src.pipeline.thinking_api":
            sub, rest = rest[:1], rest[1:]
        key = (s.env, *target, *sub)
        if key not in helps:
            helps[key] = _help(target + sub, s.env)
        known = _flags(helps[key])
        used = {a.split("=")[0] for a in rest if a.startswith("--")}
        assert used <= known, f"{s.name}: {sorted(used - known)} not in {' '.join(target + sub)} --help"


# ------------------------------------------------------------------------------------------------------------------
# money arithmetic

def test_odds_conversion_and_margin_removal():
    from src.ingest import pinnacle as P
    assert P.decimal_from_american(150) == pytest.approx(2.5)
    assert P.decimal_from_american(-200) == pytest.approx(1.5)
    assert np.isnan(P.decimal_from_american(50))
    assert L._decimal(-250) == pytest.approx(1.4) and L._decimal(2.1) == pytest.approx(2.1)
    path = L.manual_matchups([dict(tour="atp", tourney_id="2026-5014", round="R16", player_a="A", player_b="B",
                                   day="2026-10-12", odds_a=-250, odds_b=210)],
                             path=Path(os.environ.get("TMP", ".")) / "mp_manual_test.csv")
    row = pd.read_csv(path).iloc[0]
    da, db = 1.4, 3.1
    assert row["p_a"] == pytest.approx((1 / da) / (1 / da + 1 / db))
    path.unlink()


def test_curve_roots_and_reading():
    x = np.linspace(np.log(0.05 / 0.95), np.log(0.95 / 0.05), 91)
    q = 1 / (1 + np.exp(-(0.5 * x + 0.2)))               # logit q = 0.5 x + 0.2 crosses logit p = x at x = 0.4
    roots = D.curve_roots(x, q)
    assert len(roots) == 1 and roots[0] == pytest.approx(1 / (1 + np.exp(-0.4)), abs=1e-6)
    assert D.q_at(x, q, 0.5) == pytest.approx(1 / (1 + np.exp(-0.2)), abs=1e-9)
    assert np.isnan(D.q_at(x, q, 0.99))
    flat = 1 / (1 + np.exp(-(x + 1.0)))                    # always above: no crossing
    assert D.curve_roots(x, flat) == []


def test_round_mapping_from_the_draw():
    assert L._round("R1", 96) == "R128" and L._round("R2", 96) == "R64" and L._round("R3", 96) == "R32"
    assert L._round("R1", 32) == "R32" and L._round("R4", 32) == "SF" and L._round("R5", 32) == "F"
    assert L._round("R16", 96) == "R16" and L._round("QF", 96) == "QF" and L._round("Final", 32) == "F"
    assert L._round("Group A", 8) is None


def test_league_parsing():
    from src.ingest import pinnacle as P
    assert P.parse_league("ATP Shanghai - R2")["level"] == "tour"
    assert P.parse_league("ATP Challenger Villena - QF")["level"] == "challenger"
    assert P.parse_league("WTA 125k Suzhou - SF")["level"] == "125"
    assert P.parse_league("ITF Women Las Vegas - QF")["level"] == "itf"
    assert P.parse_league("ATP Shanghai - Doubles")["doubles"]
    assert P.parse_league("WTA 125k Rovereto - Qualifiers")["qualifying"]


def test_name_key_is_prediction_runs():
    from src.eda.merge_feasibility import scan as S
    for n in ("Félix Auger-Aliassime", "Alex de Minaur", "Botic van de Zandschulp", "Qinwen Zheng"):
        assert L.name_key(n) == S.name_key(n)
    assert L.name_key("Felix Auger-Aliassime") == L.name_key("Auger Aliassime Felix")


# ------------------------------------------------------------------------------------------------------------------
# Pinnacle staging and linking, on the saved probe

def _probe_pull():
    from src.ingest import pinnacle as P
    doc = {"retrieved_at": "2026-10-09T15:10:42+00:00",
           "matchups": json.loads((PROBE / "tennis_matchups.json").read_text(encoding="utf-8")),
           "straight": json.loads((PROBE / "tennis_straight.json").read_text(encoding="utf-8"))}
    return P.stage_pull(doc)


@needs_probe
def test_staging_keeps_only_pre_match_moneylines():
    t = _probe_pull()
    assert len(t) == 58 and int(t["usable"].sum()) == 46
    assert set(t.loc[~t["usable"], "reason"]) == {"in-play (live) market", "doubles"}
    u = t[t["usable"]]
    assert (u["overround"] > 1).all() and u["p_home"].between(0, 1).all()
    assert (pd.to_datetime(u["start_utc"], utc=True) > pd.Timestamp("2026-10-09T15:10:42Z")).all()


@needs_probe
def test_linker_on_the_probe(monkeypatch):
    # Tennis My Life as it stood when this test was written (11-14 Oct 2026): it held the results of many matches the
    # probe lists, but not yet Shanghai's final (18 Oct), after which every listed Shanghai round reads "a new event of
    # the same name" (found 22 Oct 2026; the qualifying test cuts the same way)
    real = L._tml
    monkeypatch.setattr(L, "_tml", lambda years: (lambda t: t[t["_date"] <= pd.Timestamp("2026-10-14")])(real(years)))
    m = L.todays_matches(pull=_probe_pull(), hours_ahead=96, now=dt.datetime(2026, 10, 9, 15, 11, tzinfo=dt.timezone.utc))
    ok = m[m["status"].eq("ok")]
    shanghai = m[m["tourney_id"].eq("2026-5014")]
    assert len(shanghai) > 10 and set(shanghai["round"].dropna()) <= {"R64", "R32"}
    assert m["reason"].astype(str).str.contains("already played").sum() >= 20          # TML holds those results now
    assert m["reason"].astype(str).str.contains("names 2 Stage 0 players").any()       # Jakub Mensik
    assert ok["stage0_id_a"].notna().all() and ok["stage0_id_b"].notna().all()
    assert ok["league"].str.match(r"^(ATP|WTA) (?!Challenger|125)").all()


# ------------------------------------------------------------------------------------------------------------------
# diagnostics on an existing replay (read only: nothing written into the daily pipeline's run folder)

@needs_replay
def test_diagnostics_on_a_replay():
    d = D.bet_diagnostics(REPLAY, write=False, show=False, now=dt.datetime(2026, 7, 15, 6, 0, tzinfo=dt.timezone.utc))
    pred = d[d["status"].eq("predicted")]
    p = pd.read_csv(REPLAY / "predictions.csv")
    assert len(pred) == len(p)
    assert np.allclose(pred["q_a"].to_numpy(), p["q_a"].round(6).to_numpy())
    assert np.allclose(pred["edge_a"], (p["q_a"] - p["p_market_a"]).round(6), atol=2e-6)
    # a replay has no bookmaker odds with their margin: no EV, no pick, anywhere (review V1, item 1)
    assert (pred["pick"] == "").all() and pred["blocked"].all()
    assert pred["block_reasons"].str.contains("no bookmaker odds").all()
    gated = pred[pred["block_reasons"].str.contains("gate:")]
    assert set(gated["gate"]) == {g for g in p["gate"] if not g.startswith("may be bet")}
    assert np.allclose(pred["fair_odds_a"], (1 / p["q_a"]).round(4), atol=1e-3)
    n = pred["break_even_p_a"].map(lambda s: len(str(s).split(";")) if isinstance(s, str) else 0)
    assert n.value_counts().to_dict() == {1: 9, 3: 7, 0: 6, 2: 1}              # review D2's count on this replay


def _american(d: float) -> float:
    return round((d - 1) * 100) if d >= 2 else round(-100 / (d - 1))


def _priced_run(tmp_path):
    """The replay as a live run made from a matchups CSV with bookmaker odds (a 3% margin): half the matches with
    Pinnacle's home player on side A, half on side B. Returns the run folder, the matchups and its rows."""
    import shutil
    run = tmp_path / "live_test_run"
    run.mkdir()
    for f in ("predictions.csv", "curves.parquet", "live_frame.parquet", "gate.csv"):
        shutil.copy2(REPLAY / f, run / f)
    shutil.copytree(REPLAY / "fit", run / "fit")
    pred = pd.read_csv(run / "predictions.csv")
    pred["player_a"] = [f"Alpha Player{i}" for i in range(len(pred))]
    pred["player_b"] = [f"Beta Player{i}" for i in range(len(pred))]
    pred.to_csv(run / "predictions.csv", index=False)
    rows, margin = [], 1.03
    for i, p in pred.iterrows():
        pa = float(p["p_market_a"])
        da, db = 1 / (pa * margin), 1 / ((1 - pa) * margin)            # side A's and B's odds with a 3% margin
        home_is_a = i % 2 == 0
        oa, ob = _american(da), _american(db)
        d_home, d_away = (L._decimal(oa), L._decimal(ob)) if home_is_a else (L._decimal(ob), L._decimal(oa))
        p_home = (1 / d_home) / (1 / d_home + 1 / d_away)
        rows.append({"tour": p["tour"], "tourney_id": "x", "round": p["round"], "day": p["day"],
                     "player_a": p["player_a"] if home_is_a else p["player_b"],
                     "player_b": p["player_b"] if home_is_a else p["player_a"],
                     "stage0_id_a": p["player_a_id"] if home_is_a else p["player_b_id"],
                     "stage0_id_b": p["player_b_id"] if home_is_a else p["player_a_id"],
                     "p_a": p_home, "odds_a": oa if home_is_a else ob,
                     "odds_b": ob if home_is_a else oa, "overround": 1 / d_home + 1 / d_away,
                     "quoted_at": "2026-07-15T05:55:00Z", "start_utc": "2026-07-15T09:00:00Z",
                     "price_source": "pinnacle"})
    mu = pd.DataFrame(rows)
    # the run was fed the CSV's prices: make predictions.csv's p_market_a exactly side A's of the CSV
    pred["p_market_a"] = [r["p_a"] if i % 2 == 0 else 1 - r["p_a"] for i, r in enumerate(rows)]
    pred.to_csv(run / "predictions.csv", index=False)
    mu.to_csv(run / "matchups_input.csv", index=False)
    return run, mu, rows


PRICED_NOW = dt.datetime(2026, 7, 15, 6, 0, tzinfo=dt.timezone.utc)      # five minutes after _priced_run's quotes


@needs_replay
def test_diagnostics_orient_the_matchups_by_id(tmp_path):
    """A run made from a matchups CSV: half the matches with Pinnacle's home player on side A, half on side B. The
    odds shown must be side A's in both cases, the price must read back, and a corrupted price or a swapped player
    must block the match."""
    run, mu, rows = _priced_run(tmp_path)
    now = PRICED_NOW
    d = D.bet_diagnostics(run, write=False, show=False, now=now)
    d = d[d["status"].eq("predicted")].reset_index(drop=True)
    assert not d["block_reasons"].astype(str).str.contains("read-back|differ").any()
    for i, r in d.iterrows():
        side_a = rows[i]["odds_a"] if i % 2 == 0 else rows[i]["odds_b"]
        assert r["pinnacle_odds_a"] == pytest.approx(L._decimal(side_a), abs=1e-4)
        assert r["quote_age_min"] == pytest.approx(5.0)
        if r["pick"]:                                                  # a pick is +EV at the odds it shows
            assert r["ev_pinnacle_pick"] > 0
            assert r["fair_odds_pick"] == pytest.approx(1 / r["p_pick"], rel=1e-3)
    assert (d["pick"] != "").any()
    # a player new since Stage 0 carries no resolved id: not a read-back failure
    newp = mu.copy()
    newp.loc[2, "stage0_id_b"] = np.nan
    newp.to_csv(run / "matchups_input.csv", index=False)
    d1 = D.bet_diagnostics(run, write=False, show=False, now=now)
    assert "differ" not in str(d1.loc[2, "block_reasons"])
    bad = mu.copy()
    bad.loc[0, "p_a"] = bad.loc[0, "p_a"] + 0.01                      # a price that does not read back
    bad.loc[1, ["stage0_id_a"]] = "atp:999999"                         # a player that does not read back
    bad.to_csv(run / "matchups_input.csv", index=False)
    d2 = D.bet_diagnostics(run, write=False, show=False, now=now)
    d2 = d2[d2["status"].eq("predicted")].reset_index(drop=True)
    assert "price read-back mismatch" in str(d2.loc[0, "block_reasons"]) and d2.loc[0, "pick"] == ""
    assert "differ from the linked ids" in str(d2.loc[1, "block_reasons"]) and d2.loc[1, "pick"] == ""
    late = D.bet_diagnostics(run, write=False, show=False, now=now + dt.timedelta(hours=1))
    assert late[late["status"].eq("predicted")]["block_reasons"].astype(str).str.contains("old").all()


# ------------------------------------------------------------------------------------------------------------------
# the bet sheet (the owner's request of 13 Oct 2026): it re-says the record, never computes

@needs_replay
def test_bet_sheet_says_what_the_record_says(tmp_path):
    """Every BET is the record's pick at its odds and EV, with the 8 Oct rule's stake; a PASS has no +EV side at
    Pinnacle's odds; a BLOCKED row offers no player, stake or odds; bets come first; the Matches sheet lists both
    players in Pinnacle's order."""
    from src.eda.tabpfn_thinking58.common import stake
    run, _, rows = _priced_run(tmp_path)
    rec = D.bet_diagnostics(run, write=False, show=False, now=PRICED_NOW, stake_rule="owner_8oct")
    pred = rec[rec["status"].eq("predicted")]
    assert set(pred["home_side"]) == {"A", "B"}
    sh = BS.build(rec, made_at=PRICED_NOW, stake_rule="owner_8oct", meta={"stale_minutes": 30})
    t = sh.table
    assert len(t) == len(pred)
    order = t["Decision"].map(BS.ORDER).to_list()
    assert order == sorted(order)                                       # BET, then PASS, then BLOCKED
    assert (t["Decision"] == "BET").sum() == pred["pick"].isin(["A", "B"]).sum() > 0
    picks = pred[pred["pick"].isin(["A", "B"])].set_index("pick_player")
    for _, r in t[t["Decision"].eq("BET")].iterrows():
        p = picks.loc[r["Player"]]
        assert r["US odds"] == BS.american(p["pinnacle_odds_pick"])
        assert r["Expected profit"] == pytest.approx(p["ev_pinnacle_pick"]) and r["Expected profit"] > 0
        assert r["Lowest odds to take"] == BS.american(1 / p["p_pick"])
        assert r["Model chance"] == pytest.approx(p["p_pick"]) and r["US odds"] >= r["Lowest odds to take"]
        assert r["Stake (units)"] == pytest.approx(float(stake(p["pinnacle_odds_pick"])), abs=1e-3)
    for _, r in t[t["Decision"].eq("PASS") & t["Player"].notna()].iterrows():
        assert r["Expected profit"] <= 0 and r["US odds"] <= r["Lowest odds to take"] and r["Edge (pts)"] > 0
    blocked = t[t["Decision"].eq("BLOCKED")]
    assert blocked[["Player", "Stake (units)", "US odds"]].isna().all().all() and blocked["Notes"].str.len().gt(0).all()
    # US odds only (the owner, 20 Oct 2026): no decimal odds column, and no decimal price in any note or header line
    every = [*t.columns, *sh.matches.columns, *sh.left_out.columns]
    assert not [c for c in every if re.fullmatch(r"(Odds|Fair odds) ?\d?", c)]
    assert {k for c, k in sh.kinds.items() if "odds" in c.lower()} == {"us"}
    assert {k for c, k in sh.matches_kinds.items() if "odds" in c.lower()} == {"us"}
    texts = [*t["Notes"].astype(str), *sh.matches["Edge gone if the market reaches"].astype(str), *sh.header]
    assert not [x for x in texts if re.search(r"\b\d+\.\d{3}\b", x)]          # decimal odds were said as 2.345
    m = sh.matches.set_index("Match id")
    home_p = {pred.iloc[i]["match_id"]: r["p_a"] for i, r in enumerate(rows)}         # the CSV's home player is 1
    assert np.allclose([m.loc[k, "Market chance 1"] for k in home_p], list(home_p.values()), atol=1e-6)
    assert np.allclose(m["Model chance 1"] + m["Model chance 2"], 1)


def test_plain_reasons():
    """The record's reasons said plainly; an unknown one kept as it is."""
    assert BS.american(1.3145) == -318 and BS.american(2.32) == 132 and BS.american(2.0) == 100
    assert np.isnan(BS.american(1.0)) and np.isnan(BS.american(None))
    gate = {"block_reasons": "gate: do not bet: inputs incomplete (1 earlier match(es) without a box score)",
            "gate_detail": "Casper Ruud's Laver Cup RR match (Tennis My Life has it, but our consistency check rejected it)"}
    assert BS.why_blocked(gate).startswith("Data gap: no usable box score for Casper Ruud's Laver Cup RR match (Tennis "
                                           "My Life has it, but our consistency check rejected it)")   # a run before 14 Oct
    since = BS.why_blocked({"block_reasons": "gate: do not bet: inputs incomplete (Jane Doe: entered through qualifying, "
                                             "but no qualifying match in Stage 0 | indoor or outdoor unknown "
                                             "(src/eda/indoor_fill/status.py found no source))"})
    assert since.startswith("Data check: Jane Doe came through qualifying") and "indoors is unknown" in since
    assert [BS.data_check_word(g) for g in ("may be bet", "may be bet if its edge clears 1 point",
                                            "do not bet: inputs incomplete")] == \
        ["complete", "gap: 1-point edge needed", "incomplete"]
    stale = BS.why_blocked({"block_reasons": "quote 45 min old (over 30): reprice; the match has started"})
    assert "45 min old (limit 30)" in stale and "A6" in stale and "has started" in stale
    assert BS.why_blocked({"block_reasons": "something new"}) == "Something new."
    assert BS.why_left_out("no main-tour Tennis My Life event named like 'Wuhan' in the last 16 days (a first round "
                           "waits for the event's first played match)").startswith("First round of a new event")
    assert BS.why_left_out("'Jakub Mensik' names 2 Stage 0 players (atp:210150, atp:210084): the run would stop"
                           ).startswith("Jakub Mensik's name matches 2 players")
    seeds = BS.match_warnings({"warnings": "the 16 seeds disagree (sd 0.034)", "player_a": "X", "player_b": "Y"})
    assert seeds and "±3.4 pts" in seeds[0]


@needs_replay
def test_bet_sheet_files(tmp_path, monkeypatch):
    """A live run's sheet goes on top of predictions/<day>/, its record and provenance into records/, and a second
    call never replaces the first."""
    from openpyxl import load_workbook
    run, _, _ = _priced_run(tmp_path)
    monkeypatch.setattr(CFG, "PREDICTIONS", tmp_path / "predictions")
    rec = D.bet_diagnostics(run, write=True, show=False, now=PRICED_NOW, stake_rule="owner_8oct")
    D.bet_diagnostics(run, write=True, show=False, now=PRICED_NOW, stake_rule="owner_8oct")
    folder = tmp_path / "predictions" / str(pd.read_csv(run / "predictions.csv")["day"].min())[:10]
    top = sorted(p.name for p in folder.iterdir() if p.is_file())
    stem = top[0][:-len(".xlsx")]
    assert top == [f"{stem}.xlsx", f"{stem}_2.xlsx"]
    assert sorted(p.name for p in (folder / "records").iterdir()) == sorted(
        [f"{stem}.csv", f"{stem}.provenance.json", f"{stem}_2.csv", f"{stem}_2.provenance.json"])
    back = pd.read_csv(folder / "records" / f"{stem}.csv")
    assert list(back.columns) == list(rec.columns) and len(back) == len(rec)
    wb = load_workbook(folder / f"{stem}.xlsx")
    assert wb.sheetnames == ["Bet sheet", "Matches", "Not predicted", "Run", "How to read"]
    cells = {str(c.value) for row in wb["Bet sheet"].iter_rows() for c in row if c.value is not None}
    assert set(rec.loc[rec["pick"].isin(["A", "B"]), "pick_player"]) <= cells
    prov = json.loads((folder / "records" / f"{stem}.provenance.json").read_text(encoding="utf-8"))
    assert prov["files"]["Bet sheet"].endswith(f"{stem}.xlsx") and "run_notes" in prov


# ------------------------------------------------------------------------------------------------------------------
# the data check (the owner's request of 13 Oct 2026): the gate's rule before the run

LIVE_RUN = ROOT / "data" / "studies" / "daily_prediction" / "runs" / "live_2026-10-11_110150"


def test_no_module_shadows_an_exported_function():
    """Importing a submodule sets it on the package: one named like an exported function would replace the function
    (tp.data_check was the module for a moment on 13 Oct)."""
    import importlib
    for mod in tp._MODULES:
        importlib.import_module(f"src.pipeline.{mod}")
    assert not set(tp._MODULES) & set(tp._NAMES)
    assert all(callable(getattr(tp, name)) and not isinstance(getattr(tp, name), type(tp)) for name in tp._NAMES)


def test_data_check_names_each_cause_and_fix(tmp_path):
    """Five matches on a synthetic Stage 0 (the rule of 14 Oct 2026): complete; a main-tour box score our checks
    rejected (a bet needs a 1-point edge); a WTA box score the run fills from the WTA's feed; a qualifier with no
    qualifying match (blocked); a team-event box score missing (noted only, ready)."""
    from src.pipeline import readiness as DCK
    s0 = tmp_path / "stage0"
    s0.mkdir()
    matches, players, timing = [], [], []
    lf = (pd.Timestamp("2026-10-16") - pd.Timestamp("1970-01-01")).days
    for tour, tid, pid, opp, status, entry, level in (
            ("atp", "2026-9001", "atp:1", "atp:91", "full", None, "T500"),
            ("atp", "2026-9001", "atp:2", "atp:92", "full", None, "T500"),
            ("atp", "2026-9001", "atp:3", "atp:93", "failed_check", None, "T500"),
            ("atp", "2026-9001", "atp:4", "atp:94", "full", None, "T500"),
            ("wta", "2026-9002", "wta:5", "wta:95", "blank", None, "T500"),
            ("wta", "2026-9002", "wta:6", "wta:96", "full", None, "T500"),
            ("wta", "2026-9002", "wta:7", "wta:97", "full", "Q", "T500"),
            ("wta", "2026-9002", "wta:8", "wta:98", "full", None, "T500"),
            ("atp", "2026-9003", "atp:10", "atp:90", "blank", None, "TEAM"),
            ("atp", "2026-9001", "atp:11", "atp:89", "full", None, "T500")):
        mid = f"{tour}|{tid}-20261014|R32|{pid}|{opp}"
        matches.append(dict(match_id=mid, tour=tour, year=2026, level=level, qualifying=False, round="R32",
                            round_order=13, tournament_id=f"{tid}-20261014", tournament_name="Testopolis",
                            tournament_date=pd.Timestamp("2026-10-14"), status="completed",
                            has_box=status == "full", box_status=status, source="tml"))
        players += [dict(match_id=mid, player_id=pid, player_name=pid, entry=entry),
                    dict(match_id=mid, player_id=opp, player_name=opp, entry=None)]
        timing.append(dict(match_id=mid, lf=lf))
    pd.DataFrame(matches).to_parquet(s0 / "stage0_matches.parquet")
    pd.DataFrame(players).to_parquet(s0 / "stage0_match_players.parquet")
    pd.DataFrame(timing).to_parquet(tmp_path / "timing.parquet")
    mu = pd.DataFrame([dict(tour=t, tourney_id=tid, round="R16", player_a=a, player_b=b, day="2026-10-20",
                            tourney_name="Testopolis", start_utc="2026-10-20T10:00:00Z", stage0_id_a=a, stage0_id_b=b)
                       for t, tid, a, b in (("atp", "2026-9001", "atp:1", "atp:2"), ("atp", "2026-9001", "atp:3", "atp:4"),
                                            ("wta", "2026-9002", "wta:5", "wta:6"),
                                            ("wta", "2026-9002", "wta:7", "wta:8"),
                                            ("atp", "2026-9001", "atp:10", "atp:11"))])
    mu.to_csv(tmp_path / "matchups.csv", index=False)
    out = DCK.data_check(tmp_path / "matchups.csv", stage0=s0, timing=tmp_path / "timing.parquet", held=set(),
                         packages=False, now=dt.datetime(2026, 10, 19, 12, 0, tzinfo=dt.timezone.utc), write=False,
                         show=False)
    got = dict(zip(out["Match"], out["Verdict"]))
    assert got == {"atp:1 v atp:2": "ready", "atp:3 v atp:4": "needs a 1-point edge",
                   "wta:5 v wta:6": "ready once filled", "wta:7 v wta:8": "will be blocked", "atp:10 v atp:11": "ready"}
    why = dict(zip(out["Match"], out["Why"] + " | " + out["What fixes it"]))
    assert "consistency check rejected it" in why["atp:3 v atp:4"] and "data repair" in why["atp:3 v atp:4"]
    assert "the WTA's own feed" in why["wta:5 v wta:6"]
    assert "came through qualifying" in why["wta:7 v wta:8"]
    noted = dict(zip(out["Match"], out["Noted only"]))
    assert "team event" in noted["atp:10 v atp:11"] and not why["atp:10 v atp:11"].strip(" |")
    assert list(out["Verdict"]) == sorted(out["Verdict"], key=lambda v: -list(DCK.VERDICTS).index(v))   # blocks first


@pytest.mark.skipif(not (LIVE_RUN / "stage0" / "stage0_matches.parquet").exists(), reason="no live run of 11 Oct")
def test_data_check_agrees_with_the_runs_own_gate(tmp_path):
    """On the 11 Oct run's own Stage 0 (the builder before the match-tiebreak repair) and matchups, the data check
    checks the same earlier matches as the run's gate, and finds the gaps the run's gate found. Under the rule of 14 Oct
    2026 the three Laver Cup gaps that blocked three matches that day are a team event's: noted only, nothing blocked."""
    from src.pipeline import readiness as DCK
    # since the shared rebuild of 13 Oct 21:13-21:16 the shared timeline dates lower-level matches that this older
    # Stage 0 lacks and the feed adds again, and the merge refuses the duplicate timeline rows (the same collision that
    # stops the engine's replays; box_score_gate_review, process log Part 1, 22:15-22:35): the timeline is restricted
    # to this Stage 0's matches, as that Stage 0's own run had it
    ids = set(pd.read_parquet(LIVE_RUN / "stage0" / "stage0_matches.parquet", columns=["match_id"])["match_id"].astype(str))
    t = pd.read_parquet(CFG.TIMING)
    t[t["match_id"].astype(str).isin(ids)].to_parquet(tmp_path / "timing.parquet", index=False)
    out = DCK.data_check(LIVE_RUN / "matchups_input.csv", stage0=LIVE_RUN / "stage0", timing=tmp_path / "timing.parquet",
                         now=dt.datetime(2026, 10, 11, 18, 1, tzinfo=dt.timezone.utc), write=False, show=False)
    gate = pd.read_csv(LIVE_RUN / "gate.csv").merge(pd.read_csv(LIVE_RUN / "predictions.csv")[
        ["match_id", "player_a", "player_b"]], on="match_id")
    from src.eda.merge_feasibility.scan import name_key          # the CSV spells names as Stage 0 does, the run too
    want = {frozenset(name_key(x) for x in (r.player_a, r.player_b)): (r.without_box, r.earlier_matches_checked)
            for r in gate.itertuples()}
    got = {frozenset(name_key(x) for x in m.split(" v ")): (a + b, n)
           for m, a, b, n in zip(out["Match"], out["Missing: edge needed"], out["Missing: noted only"],
                                 out["Earlier matches checked"])}
    assert got == want and sum(g for g, _ in got.values()) == 3
    assert set(out["Verdict"]) == {"ready"} and out["Missing: noted only"].sum() == 3
    assert out["Noted only"].str.contains("Laver Cup RR match \\(team event").sum() == 3


@needs_replay
def test_the_bet_sheet_applies_the_edge_bar(tmp_path):
    """A match whose gate sets the 1-point bar is bet only when its forecast, moved 1 point against the side, is still
    +EV at the odds (the rule the gate review's replay tested); a +EV side that does not clear it is a PASS that says
    so, and nothing is blocked."""
    run, _, _ = _priced_run(tmp_path)
    pred = pd.read_csv(run / "predictions.csv")
    pred["gate"], pred["gate_reason"], pred["gate_edge_bar"] = (
        "may be bet if its edge clears 1 point", "1 earlier match(es) without a box score of a kind the training rows "
        "nearly always hold: the edge must clear 1 point", 0.01)
    pred.to_csv(run / "predictions.csv", index=False)
    rec = D.bet_diagnostics(run, write=False, show=False, now=PRICED_NOW)
    p = rec[rec["status"].eq("predicted")]
    assert not p["blocked"].any() and (p["gate_edge_bar"] == 0.01).all()
    for _, r in p.iterrows():
        qa, da, db = r["q_a"], r["pinnacle_odds_a"], r["pinnacle_odds_b"]
        clears = {"A": (qa - 0.01) * da - 1 > 0, "B": (1 - qa - 0.01) * db - 1 > 0}
        want = "A" if clears["A"] else ("B" if clears["B"] else "")
        assert r["pick"] == want
        plain = "A" if r["ev_a_pinnacle"] > 0 else ("B" if r["ev_b_pinnacle"] > 0 else "")
        assert (r["passed_by_bar"] if isinstance(r["passed_by_bar"], str) else "") == ("" if want else plain)
    barred = p[p["passed_by_bar"].isin(["A", "B"])]
    assert len(barred) and (barred["decision"] == "PASS").all()
    assert all("Passed for a data gap" in BS.notes_for(r) for _, r in barred.iterrows())


def _log_odds_mean(a, b):
    la, lb = np.log(np.asarray(a) / (1 - np.asarray(a))), np.log(np.asarray(b) / (1 - np.asarray(b)))
    return 1 / (1 + np.exp(-(0.5 * la + 0.5 * lb)))


def _thinking_run(tmp_path, ok: bool = True):
    """The priced run as a run with TabPFN-3.5-Thinking beside the local model (what dp_predict writes): q_a is the
    ensemble's (the two chances averaged 50/50 in log-odds), the local and Thinking chances beside it; or, ok=False, a
    run whose Thinking call failed."""
    run, mu, rows = _priced_run(tmp_path)
    pred = pd.read_csv(run / "predictions.csv")
    local = pred["q_a"].to_numpy()
    thinking = np.clip(local + np.where(np.arange(len(pred)) % 3 == 0, 0.06, -0.04), 0.02, 0.98)
    pred["q_a_local"], pred["q_a_local_patterns_off"] = local, local
    pred["q_a_thinking"] = thinking if ok else np.nan
    pred["q_a"] = pred["q_a_patterns_off"] = _log_odds_mean(local, thinking) if ok else local
    pred["patterns"], pred["ensemble"], pred["ensemble_weight_local"] = "off", "on" if ok else "off", 0.5 if ok else np.nan
    pred.to_csv(run / "predictions.csv", index=False)
    th = {"requested": True, "ensemble": "on", "weight_local": 0.5,
          "seasons": {"2026": {"fitted_train_set_id": "fit-x", "query_rows": 92, "tokens_used_daily_counter": 31572,
                               "served": {"weights": "tabpfn-v3.5.safetensors", "package_version": "9.0.0"}}}} if ok \
        else {"requested": True, "ensemble": "off", "failed": "thinking failed: the API said 429"}
    (run / "provenance.json").write_text(json.dumps({"thinking": th}), encoding="utf-8")
    (run / "pipeline_model.json").write_text(json.dumps({"model": "tabpfn35_thinking"}), encoding="utf-8")
    return run, local, thinking


@needs_replay
def test_the_bet_sheet_shows_both_models_and_decides_on_their_average(tmp_path):
    """A run with Thinking (the owner, 20 Oct 2026): the record carries each model's chance, the decision, edge and
    expected profit are the ensemble's, and the sheet shows the non-thinking, the thinking and the ensemble chance
    without patterns where a local run shows one column (and no decimal odds: net +1 column)."""
    run, local, thinking = _thinking_run(tmp_path)
    rec = D.bet_diagnostics(run, write=False, show=False, now=PRICED_NOW, stake_rule="owner_8oct")
    p = rec[rec["status"].eq("predicted")].reset_index(drop=True)
    assert p["model"].eq("tabpfn35_thinking").all()
    assert np.allclose(p["q_a"], _log_odds_mean(local, thinking).round(6), atol=1e-6)
    assert np.allclose(p["q_a_local_patterns_off"], local.round(6)) and np.allclose(p["q_a_thinking"], thinking.round(6))
    assert np.allclose(p["edge_a"], (p["q_a"] - p["p_market_a"]).round(6), atol=2e-6)
    for _, r in p[p["pick"].isin(["A", "B"])].iterrows():                  # each pick is +EV for the ensemble
        q = r["q_a"] if r["pick"] == "A" else 1 - r["q_a"]
        assert r["ev_pinnacle_pick"] > 0 and q * r["pinnacle_odds_pick"] - 1 == pytest.approx(r["ev_pinnacle_pick"],
                                                                                               abs=1e-4)
    sh = BS.build(rec, made_at=PRICED_NOW, stake_rule="owner_8oct", meta={"stale_minutes": 30,
                                                                          "model": "tabpfn35_thinking"})
    (tmp_path / "local").mkdir()
    local_sheet = BS.build(D.bet_diagnostics(_priced_run(tmp_path / "local")[0], write=False, show=False,
                                             now=PRICED_NOW, stake_rule="owner_8oct"),
                           made_at=PRICED_NOW, stake_rule="owner_8oct", meta={"stale_minutes": 30})
    t = sh.table
    assert list(t.columns[t.columns.get_loc("Model chance"):t.columns.get_loc("Market chance")]) == \
        ["Model chance", "Non-thinking chance without patterns", "Thinking chance without patterns",
         "Ensemble chance without patterns"]
    assert len(t.columns) == len(local_sheet.table.columns) + 2
    assert "Model chance without patterns" in local_sheet.table.columns and "Odds" not in local_sheet.table.columns
    shown = t[t["Player"].notna()]
    assert len(shown) and np.allclose(shown["Ensemble chance without patterns"],
                                      _log_odds_mean(shown["Non-thinking chance without patterns"],
                                                     shown["Thinking chance without patterns"]), atol=2e-6)
    assert np.allclose(shown["Model chance"], shown["Ensemble chance without patterns"])     # patterns off here
    assert {"Non-thinking chance 1 without patterns", "Thinking chance 1 without patterns",
            "Ensemble chance 1 without patterns"} <= set(sh.matches.columns)
    assert any("log-odds" in h for h in sh.header) and not any("plain" in h for h in sh.header)


@needs_replay
def test_a_failed_thinking_call_gives_the_local_sheet_and_says_so(tmp_path, monkeypatch):
    """Thinking asked for but its call failed: the record and the sheet are the local model's alone, the sheet's run
    notes say why (once), and predict_day's line says it loudly."""
    run, local, _ = _thinking_run(tmp_path, ok=False)
    monkeypatch.setattr(CFG, "PREDICTIONS", tmp_path / "predictions")
    rec = D.bet_diagnostics(run, write=True, show=False, now=PRICED_NOW)
    p = rec[rec["status"].eq("predicted")]
    assert p["model"].eq("tabpfn35").all() and "q_a_thinking" not in rec.columns
    assert np.allclose(p["q_a"], local.round(6))
    prov = json.loads(next((tmp_path / "predictions").rglob("*.provenance.json")).read_text(encoding="utf-8"))
    first = prov["run_notes"][0]
    assert "Thinking was asked for, but its call failed" in first["plain"] and "429" in first["detail"]
    from src.pipeline import predict as PD
    line = PD.thinking_outcome(run)
    assert line.startswith("WARNING") and "429" in line and "ALONE" in line


# ------------------------------------------------------------------------------------------------------------------
# Thinking: nothing is ever paid without the flag; the hash recipe is the daily pipeline's

def test_thinking_fit_refused_without_the_flag():
    out = subprocess.run([str(CFG.interpreter("api")), "-m", "src.pipeline.thinking_api", "fit", "--folder", "nowhere"],
                         cwd=ROOT, capture_output=True, text=True, encoding="utf-8", errors="replace")
    assert out.returncode == 2 and "Refused" in out.stdout


def test_the_daily_pipeline_offers_thinking_and_a_pipeline_without_it_is_refused(monkeypatch):
    assert tp.predict.dp_takes_thinking()                      # run_day --thinking --spend-tokens (20 Oct 2026)
    monkeypatch.setattr("src.pipeline.predict.dp_takes_thinking", lambda: False)
    with pytest.raises(NotImplementedError):
        tp.predict_day("x.csv", model="tabpfn35_thinking", spend_tokens=True)


def test_predict_day_with_thinking_makes_a_missing_fit_first_then_runs_both(tmp_path, monkeypatch):
    """A dry run: with no season fit, the paid fit's two steps come first, then run_day with --thinking and
    --spend-tokens; with the fit, only the run; a fit made on another frame is refused. Every flag parses."""
    from src.pipeline import thinking as TH
    mu = tmp_path / "m.csv"
    pd.DataFrame({"day": ["2026-10-21"]}).to_csv(mu, index=False)
    seen: list[R.Step] = []

    def capture(step, **kw):
        seen.append(step)
        return R.StepResult(step.name, step.shown(), "dry-run", None, 0.0, None)
    monkeypatch.setattr("src.pipeline.predict.run_step", capture)
    monkeypatch.setattr("src.pipeline.runner.run_steps", lambda steps, **kw: [capture(s) for s in steps])
    monkeypatch.setattr(TH, "fit_state", lambda instance, season: "missing")
    assert tp.predict_day(mu, model="tabpfn35_thinking", spend_tokens=True, skip_checks=True, dry_run=True) is None
    assert [s.name for s in seen] == ["thinking_prepare_2026", "thinking_fit_2026", "predict_day"]
    assert "--spend-tokens" in seen[1].args and seen[1].env == "api"
    run_day = [str(a) for a in seen[2].args]
    assert {"--thinking", "--spend-tokens", "--refresh"} <= set(run_day)
    known = _flags(_help(run_day[:2]))
    assert {a for a in run_day if a.startswith("--")} <= known
    seen.clear()
    monkeypatch.setattr(TH, "fit_state", lambda instance, season: "ok")
    tp.predict_day(mu, model="tabpfn35_thinking", spend_tokens=True, skip_checks=True, dry_run=True)
    assert [s.name for s in seen] == ["predict_day"]
    monkeypatch.setattr(TH, "fit_state", lambda instance, season: "other frame")
    with pytest.raises(RuntimeError, match="another frame"):
        tp.predict_day(mu, model="tabpfn35_thinking", spend_tokens=True, skip_checks=True, dry_run=True)
    seen.clear()
    tp.predict_day(mu, model="tabpfn35", skip_checks=True, dry_run=True)       # the local model: no Thinking at all
    assert [s.name for s in seen] == ["predict_day"] and "--thinking" not in [str(a) for a in seen[0].args]


def test_fit_state_reads_the_instance_frame(tmp_path, monkeypatch):
    from src.pipeline import thinking as TH
    monkeypatch.setattr(CFG, "THINKING", tmp_path / "thinking")
    inst = tmp_path / "template_x"
    inst.mkdir()
    (inst / "model_frame.parquet").write_bytes(b"frame one")
    assert TH.fit_state(inst, 2026) == "missing"
    out = TH.folder(inst, 2026)
    out.mkdir(parents=True)
    import hashlib
    for n in TH.FIT_FILES:
        (out / n).write_text("{}", encoding="utf-8")
    (out / "prepared.json").write_text(json.dumps({"frame_sha256": hashlib.sha256(b"frame one").hexdigest()}),
                                       encoding="utf-8")
    assert TH.fit_state(inst, 2026) == "ok"
    (inst / "model_frame.parquet").write_bytes(b"frame two")
    assert TH.fit_state(inst, 2026) == "other frame"


def test_train_sha256_matches_the_daily_pipelines_recipe():
    sys.path.insert(0, str(CFG.DAILY))
    import importlib
    tc = importlib.import_module("dp_tabcache")
    X = pd.DataFrame({"a": [1.0, np.nan, 3.0], "b": [0.5, 0.25, -1.0]})
    y = np.array([1, 0, 1])
    from src.pipeline.thinking import train_sha256
    assert train_sha256(X.columns, X.to_numpy(), y) == tc.train_sha256(X, y)


# ------------------------------------------------------------------------------------------------------------------
# the notebook

def test_notebook_is_clean_and_calls_only_existing_functions():
    import nbformat
    nb = nbformat.read(str(NOTEBOOK), as_version=4)
    nbformat.validate(nb)
    names = set()
    for c in nb.cells:
        if c.cell_type != "code":
            continue
        assert not c.get("outputs") and c.get("execution_count") is None
        compile(c.source, "<cell>", "exec")
        names |= set(re.findall(r"\btp\.([a-z_]+)", c.source))
    missing = [n for n in names if not hasattr(tp, n)]
    assert not missing, missing
    first = next(c for c in nb.cells if c.cell_type == "code").source
    assert "MODEL = " in first and "DRY_RUN = " in first


def test_reprice_dry_run_pulls_nothing(monkeypatch):
    monkeypatch.setattr("src.pipeline.live.record_pinnacle", lambda **k: (_ for _ in ()).throw(AssertionError("pulled")))
    assert tp.reprice(REPLAY, dry_run=True) is None


def test_latest_run_is_by_name(tmp_path, monkeypatch):
    monkeypatch.setattr(CFG, "DAILY_RUNS", tmp_path)
    old, new = tmp_path / "live_2026-10-11_090000", tmp_path / "live_2026-10-11_100000"
    new.mkdir(); old.mkdir()                                     # the older run was touched last
    (old / "reprice_x.csv").write_text("x")
    assert tp.latest_run() == new


def test_thinking_needs_spend_tokens(monkeypatch):
    monkeypatch.setattr("src.pipeline.predict.dp_takes_thinking", lambda: True)
    with pytest.raises(RuntimeError, match="spend_tokens"):
        tp.predict_day("x.csv", model="tabpfn35_thinking")


def test_event_over_or_stale_is_refused():
    rows = pd.DataFrame({"round": ["R32", "R16", "QF", "SF", "F"],
                         "_date": pd.to_datetime(["2023-01-02", "2023-01-03", "2023-01-04", "2023-01-06", "2023-01-07"])})
    assert "final is played" in L._event_current(rows, "R32", pd.Timestamp("2023-01-09"))
    ongoing = rows.iloc[:2]
    assert L._event_current(ongoing, "R16", pd.Timestamp("2023-01-04")) is None
    assert "over a week" in L._event_current(ongoing, "R16", pd.Timestamp("2023-01-20"))


def test_manual_rows_need_time_zones(tmp_path):
    with pytest.raises(ValueError, match="time zone"):
        L.manual_matchups([dict(tour="atp", tourney_id="2026-5014", round="R16", player_a="A", player_b="B",
                                day="2026-10-12", odds_a=-150, odds_b=130, quoted_at="2026-10-12T08:00:00")],
                          path=tmp_path / "m.csv")


def test_dp_env_hands_over_the_checked_tree(monkeypatch):
    monkeypatch.setenv("DP_BASE_TREE", "C:/somewhere/else")      # never read from the kernel
    env = tp.train.dp_env()
    assert env["DP_BASE_TREE"] == str(tp.train.base_tree())
    assert "somewhere" not in env["DP_BASE_TREE"]
