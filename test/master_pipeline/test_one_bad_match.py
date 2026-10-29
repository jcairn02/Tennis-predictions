"""One bad row never holds up the others, and is said loudly (the owner, 28 Oct 2026: "while you fix a single row it
shouldnt hold up the whole pipeline", and "Be sure to mention the match that was left out and that there was an error
in one row clearly and loudly, however still give all the other predictions"; docs/decisions/2026-10-28-one-bad-match-
left-out.md; master pipeline process log, Part 25): A2 (one listing), A3 (one matchups row), A3b (one match checked),
A4 (the rows its check flags, and the engine's own), A5 (the sheet's alerts and its NOT PREDICTED rows), A5b (one
match repriced), A7 (one run graded), A8 (one sheet's record)."""
from __future__ import annotations

import datetime as dt
import json
import shutil
import sys
from pathlib import Path

import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import src.pipeline as tp  # noqa: E402
from src.eda.merge_feasibility.scan import name_key  # noqa: E402
from src.pipeline import bet_sheet as BS  # noqa: E402
from src.pipeline import closing_line as CL  # noqa: E402
from src.pipeline import config as CFG  # noqa: E402
from src.pipeline import diagnostics as D  # noqa: E402
from src.pipeline import left_out as LO  # noqa: E402
from src.pipeline import live as L  # noqa: E402
from src.pipeline import predict as PR  # noqa: E402
from src.pipeline import runner as R  # noqa: E402

NOW = pd.Timestamp("2026-10-28T12:00:00Z")
REPLAY = ROOT / "data" / "studies" / "daily_prediction" / "runs" / "replay_2026-07-13_R16_speed"
needs_replay = pytest.mark.skipif(not (REPLAY / "predictions.csv").exists(), reason="no replay run")


@pytest.fixture(autouse=True)
def men_only(monkeypatch):
    monkeypatch.setattr(CFG, "BET_TOURS", ("atp",))


def _pull(leagues: list[str]) -> pd.DataFrame:
    """A staged Pinnacle pull: one pre-match moneyline per league name, two hours ahead."""
    from src.ingest import pinnacle as P
    t = (NOW + pd.Timedelta(hours=2)).strftime("%Y-%m-%dT%H:%M:%SZ")
    matchups = [{"id": i, "parentId": None, "isLive": False, "status": "pending", "startTime": t,
                 "league": {"name": name, "id": i, "bestOfX": 3},
                 "participants": [{"alignment": "home", "name": f"Home Player{i}"},
                                  {"alignment": "away", "name": f"Away Player{i}"}]}
                for i, name in enumerate(leagues, start=1)]
    straight = [{"type": "moneyline", "period": 0, "matchupId": i, "status": "open", "isAlternate": False,
                 "cutoffAt": t, "prices": [{"designation": "home", "price": -150}, {"designation": "away", "price": 130}],
                 "limits": [{"type": "maxRiskStake", "amount": 500}], "version": 1}
                for i in range(1, len(leagues) + 1)]
    return P.stage_pull({"retrieved_at": NOW.isoformat(), "matchups": matchups, "straight": straight})


# --- A2 ---------------------------------------------------------------------------------------------------------------
class _NoTennisLiveEvent:
    stopped = None

    def event(self, tour, event, day):
        return None, "no TennisLive event in this test"


def test_a2_an_error_linking_one_listing_leaves_it_out_and_links_the_others(monkeypatch, capsys):
    tml_cols = ["tour", "_tid", "_wk", "_lk", "round", "_date", "tourney_name", "winner_name", "loser_name"]
    monkeypatch.setattr(L, "_tml", lambda years: pd.DataFrame(columns=tml_cols))
    monkeypatch.setattr(L, "_stage0_keys", lambda tour: pd.DataFrame({
        "player_id": ["atp:1", "atp:2"], "player_name": ["Home Player1", "Away Player1"],
        "key": [name_key("Home Player1"), name_key("Away Player1")]}))
    monkeypatch.setattr(L, "_events", lambda tml, tour, day: pd.DataFrame())
    good = pd.Series({"_tid": "2026-9", "tourney_name": "Goodville Open", "tourney_level": "250", "draw_size": 32,
                      "best_of": 3, "last": pd.Timestamp("2026-10-27")})

    def match_event(event, cands):
        if event == "Boomtown":
            raise KeyError("a broken listing")
        return good, None
    monkeypatch.setattr(L, "_match_event", match_event)
    m = L.todays_matches(pull=_pull(["ATP Goodville - R16", "ATP Boomtown - R16"]), hours_ahead=36,
                         now=NOW.to_pydatetime(), tennislive=_NoTennisLiveEvent())
    by = m.set_index("home")                             # Home Player1 at Goodville, Home Player2 at Boomtown
    assert by.loc["Home Player1", "status"] == "ok" and by.loc["Home Player1", "player_a"] == "Home Player1"
    assert by.loc["Home Player2", "status"] == "left out"
    assert by.loc["Home Player2", "reason"] == f"{L.LINK_ERROR} (KeyError: 'a broken listing')"
    out = capsys.readouterr().out
    assert "!!! ERROR while linking 1 listing: left out, no prediction; the other 1 listings were linked" in out
    assert "Home Player2 v Away Player2" in out
    assert "!!! 1 of them because of an ERROR while linking" in L._off_note(m)
    assert BS.why_left_out(by.loc["Home Player2", "reason"]).startswith(BS.ERROR + " while linking it")


# --- A3 ---------------------------------------------------------------------------------------------------------------
def test_a3_flags_each_bad_row_loudly_and_raises_for_none(tmp_path, monkeypatch, capsys):
    csv = tmp_path / "matchups.csv"
    row = {"tour": "atp", "tourney_id": "2026-9", "round": "R16", "player_a": "Home Player1",
           "player_b": "Away Player1", "day": "2026-10-29", "p_a": 0.55, "tourney_name": "Goodville Open"}
    pd.DataFrame([row, {**row, "p_a": 1.5}, {**row, "p_a": "abc"}, {**row, "day": "garbage"},
                  {**row, "player_b": "Nobody Atall"}]).to_csv(csv, index=False)
    tml = pd.DataFrame({"tour": ["atp"], "_tid": ["2026-9"], "tourney_name": ["Goodville Open"],
                        "_wk": [name_key("Home Player1")], "_lk": [name_key("Away Player1")]})
    monkeypatch.setattr(L, "_tml", lambda years: tml)
    monkeypatch.setattr(L, "_stage0_keys", lambda tour: pd.DataFrame({
        "player_id": ["atp:1", "atp:2"], "player_name": ["Home Player1", "Away Player1"],
        "key": [name_key("Home Player1"), name_key("Away Player1")]}))
    res = L.check_matchups(csv)
    assert res["ok"].tolist() == [True, False, False, False, False]
    probs = res["problem"].tolist()
    assert probs[1] == "p_a must lie strictly between 0 and 1"
    assert probs[2].startswith("ERROR while checking this row (ValueError")
    assert probs[3] == "its day 'garbage' cannot be read"
    assert probs[4] == "'Nobody Atall' is unknown to Stage 0 and Tennis My Life"
    out = capsys.readouterr().out
    assert "!!! 4 OF 5 MATCHES WILL BE LEFT OUT OF THE RUN, NO PREDICTION" in out and "A4 predicts the other 1" in out
    pd.DataFrame([{"tour": "atp"}]).to_csv(csv, index=False)
    with pytest.raises(ValueError, match="columns missing"):        # a CSV without its columns: nothing can run
        L.check_matchups(csv)


# --- A4 ---------------------------------------------------------------------------------------------------------------
@pytest.fixture()
def a4(tmp_path, monkeypatch):
    """predict_day with everything outside it stood in: the instance and its checks, the work tree, the lock, and the
    engine (which reads the CSV it is given, predicts its rows and leaves one more out itself)."""
    from src.pipeline import train as T
    monkeypatch.setattr(T, "active_instance", lambda: tmp_path / "template_x")
    monkeypatch.setattr(T, "show_model_age", lambda instance: None)
    monkeypatch.setattr(T, "dp_env", lambda: {"DP_BASE_TREE": "tree", "DP_INSTANCE": "template_x"})
    monkeypatch.setattr(T, "template_check", lambda instance, season=None, tree=None: {"ok": True})
    monkeypatch.setattr(T, "base_tree", lambda: "tree")
    monkeypatch.setattr(CFG, "DAILY_RUNS", tmp_path / "runs")
    monkeypatch.setattr(PR, "WORK_TREE", tmp_path / "no_tree")
    monkeypatch.setattr(PR, "LOCK", tmp_path / "predict.lock")
    monkeypatch.setattr(PR, "gpu_users", lambda: [])
    given = {}

    def engine(step, **kw):
        args = [str(a) for a in step.args]
        run, csv = Path(args[args.index("--run") + 1]), Path(args[args.index("--matchups") + 1])
        given["csv"], given["rows"] = csv, pd.read_csv(csv)
        given["sidecar"] = csv.with_name(csv.stem + "_earlier.csv").exists()
        mu = given["rows"]
        pd.DataFrame({"match_id": [f"atp|x|R16|{i}" for i in range(len(mu) - 1)], "instance": "template_x",
                      "gate": "may be bet"}).to_csv(run / "predictions.csv", index=False)
        held = json.loads((run / LO.FILE).read_text(encoding="utf-8")) if (run / LO.FILE).exists() else {}
        last = mu.iloc[-1]                              # the engine's own: the last row cannot be placed
        held.setdefault("matches", []).append({"csv_row": len(mu) - 1, "tour": "atp", "tourney_name": "Basel",
                                               "round": "R16", "player_a": last["player_a"],
                                               "player_b": last["player_b"], "step": "tables",
                                               "reason": "Stage 0 built its live row as another id"})
        (run / LO.FILE).write_text(json.dumps(held), encoding="utf-8")
        return R.StepResult(step.name, step.shown(), "ok", 0, 0.0, None)
    monkeypatch.setattr(PR, "run_step", engine)
    return tmp_path, given


def _checked(rows):
    """check_matchups' stand-in: ok unless the row's player_b says what is wrong."""
    def check(csv):
        df = pd.read_csv(csv)
        return pd.DataFrame([{"row": i, "tour": r["tour"], "event": "Basel", "match": f"{r['player_a']} v "
                              f"{r['player_b']}", "round": "R16", "ok": r["player_b"] not in rows,
                              "problem": rows.get(r["player_b"]), "switched_off": r["tour"] == "wta"}
                             for i, r in df.iterrows()])
    return check


def test_a4_leaves_out_the_rows_its_check_flags_and_says_every_match_left_out_last(a4, monkeypatch, capsys):
    tmp, given = a4
    csv = tmp / "matchups.csv"
    pd.DataFrame([{"tour": "atp", "tourney_id": "2026-9", "round": "R16", "player_a": f"P{i}", "player_b": b,
                   "day": "2026-10-29", "p_a": 0.5, "notes": "a note, with a comma"}
                  for i, b in enumerate(["Q0", "Bad Price", "Q2", "Q3", "Woman B"])]).assign(
        tour=lambda d: ["atp", "atp", "atp", "atp", "wta"]).to_csv(csv, index=False)
    pd.DataFrame([{"tour": "atp", "round": "Q1", "winner_name": "P0", "loser_name": "Z"}]).to_csv(
        tmp / "matchups_earlier.csv", index=False)
    monkeypatch.setattr(L, "check_matchups", _checked({"Bad Price": "p_a must lie strictly between 0 and 1",
                                                       "Woman B": CFG.tour_off("wta") or "switched off"}))
    run = tp.predict_day(csv, run_name="live_2026-10-28_200000")
    assert given["csv"] == run / "matchups_checked.csv" and given["sidecar"]
    assert given["rows"]["player_b"].tolist() == ["Q0", "Q2", "Q3"]           # the flagged rows left out
    assert given["rows"]["notes"].tolist() == ["a note, with a comma"] * 3    # the rest's text as written
    assert pd.read_csv(run / "matchups_input.csv").equals(pd.read_csv(csv))   # the CSV as given, kept
    held = LO.read(run)["matches"]                       # A4's rows by the CSV as given, the engine's by its copy
    assert [(r["step"], r.get("row_given", r.get("csv_row")), bool(r.get("switched_off"))) for r in held] == [
        ("check", 1, False), ("check", 4, True), ("tables", 2, False)]
    lines = capsys.readouterr().out.strip().splitlines()
    tail = lines[-4:]
    assert tail[0] == ("!!! ERROR IN 2 ROWS: 2 MATCHES LEFT OUT, NO PREDICTION; the other 2 matches were predicted "
                       "(the bet sheet lists them under 'Not predicted'):")
    assert tail[1].startswith("!!!   ATP 2026-9 R16: P1 v Bad Price -- at A4's check before the run: p_a must")
    assert tail[2].startswith("!!!   ATP Basel R16: P3 v Q3 -- at the run's check of its own data: Stage 0 built")
    assert tail[3].startswith("1 match of a switched-off tour left out before the run")


def test_a4_with_every_row_flagged_makes_no_run(a4, monkeypatch, capsys):
    tmp, given = a4
    csv = tmp / "matchups.csv"
    pd.DataFrame([{"tour": "atp", "tourney_id": "2026-9", "round": "R16", "player_a": "P0", "player_b": "Bad",
                   "day": "2026-10-29", "p_a": 2.0}]).to_csv(csv, index=False)
    monkeypatch.setattr(L, "check_matchups", _checked({"Bad": "p_a must lie strictly between 0 and 1"}))
    assert tp.predict_day(csv, run_name="live_2026-10-28_200001") is None
    assert not given and not (tmp / "runs" / "live_2026-10-28_200001").exists()
    assert "!!! Every match of" in capsys.readouterr().out


def test_left_out_lines_say_nothing_for_a_run_that_left_nothing_out(tmp_path):
    assert LO.lines(tmp_path) == [] and LO.errors(tmp_path) == []
    (tmp_path / LO.FILE).write_text("{not json", encoding="utf-8")
    assert LO.read(tmp_path)["matches"][0]["unreadable"]                     # an unreadable record is said, too


# --- A5 ---------------------------------------------------------------------------------------------------------------
SHEET_NOW = dt.datetime(2026, 7, 15, 6, 0, tzinfo=dt.timezone.utc)


@needs_replay
def test_a5_puts_a_match_left_out_by_an_error_above_everything_in_red(monkeypatch, tmp_path):
    from openpyxl import load_workbook
    run = tmp_path / "live_2026-07-15_055500"
    run.mkdir()
    for f in ("predictions.csv", "curves.parquet", "live_frame.parquet", "gate.csv"):
        shutil.copy2(REPLAY / f, run / f)
    shutil.copytree(REPLAY / "fit", run / "fit")
    pred = pd.read_csv(run / "predictions.csv")                 # the replay's rows carry no names: give them some
    pred["player_a"] = [f"Alpha Player{i}" for i in range(len(pred))]
    pred["player_b"] = [f"Beta Player{i}" for i in range(len(pred))]
    pred.to_csv(run / "predictions.csv", index=False)
    man = pred[pred["tour"].eq("atp")].iloc[0]
    (run / LO.FILE).write_text(json.dumps({
        "matches": [{"csv_row": 3, "tour": "atp", "tourney_name": "Basel", "round": "R16", "player_a": "Zed Alpha",
                     "player_b": "Yan Beta", "day": "2026-07-15", "step": "tables",
                     "reason": "Stage 0 built its live row as another id"}],
        "earlier_rounds": [{"tour": "atp", "tourney_name": "Basel", "round": "Q2", "winner": man["player_a"],
                            "loser": "Xavier Gamma", "step": "tables", "reason": "not in the run's Stage 0"}]}),
        encoding="utf-8")
    pd.DataFrame([{"tour": "atp", "tourney_id": "2026-9", "round": "R16", "player_a": "Zed Alpha",
                   "player_b": "Yan Beta", "day": "2026-07-15", "p_a": 0.6, "odds_a": -150, "odds_b": 130,
                   "start_utc": "2026-07-15T09:00:00Z"}]).to_csv(run / "matchups_input.csv", index=False)
    monkeypatch.setattr(CFG, "PREDICTIONS", tmp_path / "predictions")
    rec = D.bet_diagnostics(run, write=True, show=False, now=SHEET_NOW)
    gone = rec[rec["status"].eq("not predicted")]
    assert gone["player_a"].tolist() == ["Zed Alpha"] and gone["block_reasons"].iloc[0].startswith(D.LEFT_OUT_ERROR)
    assert gone["pinnacle_odds_a"].iloc[0] == pytest.approx(L._decimal(-150))
    warned = rec[rec["status"].eq("predicted") & rec["warnings"].astype(str).str.contains("MISSING EARLIER ROUND")]
    assert warned["player_a"].tolist() == [man["player_a"]]
    prov = json.loads(next((tmp_path / "predictions").rglob("*.provenance.json")).read_text(encoding="utf-8"))
    alerts = prov["alerts"]
    assert alerts[0].startswith("ERROR IN 1 ROW: 1 MATCH LEFT OUT OF THIS RUN, SO NO PREDICTION")
    assert alerts[1] == ("No prediction: ATP Basel R16: Zed Alpha v Yan Beta. The error: Stage 0 built its live row as "
                         "another id (found at the run's check of its own data).")
    assert alerts[2].startswith("1 EARLIER ROUND MISSING")
    wb = load_workbook(next((tmp_path / "predictions").rglob("*.xlsx")))
    ws = wb["Bet sheet"]
    assert ws["A2"].value == alerts[0] and ws["A2"].fill.start_color.rgb.endswith(BS.ALERT[0].lstrip("#"))
    wn = wb["Not predicted"]
    assert str(wn.cell(row=4, column=7).value).startswith(BS.ERROR)          # its row first, in red
    assert wn.cell(row=4, column=1).fill.start_color.rgb.endswith(BS.ERROR_ROW[0].lstrip("#"))
    html = BS.to_html(BS.build(rec, made_at=SHEET_NOW, alerts=alerts))
    assert html.index("ERROR IN 1 ROW") < html.index("to bet")                # above everything in the notebook


# --- A3b ---------------------------------------------------------------------------------------------------------------
def test_a3b_an_error_checking_one_match_gives_it_check_failed_and_checks_the_others(tmp_path, capsys):
    from src.pipeline import readiness as DCK
    s0 = tmp_path / "stage0"
    s0.mkdir()
    mid = "atp|2026-9001-20261014|R32|atp:1|atp:91"
    pd.DataFrame([dict(match_id=mid, tour="atp", year=2026, level="T500", qualifying=False, round="R32",
                       round_order=13, tournament_id="2026-9001-20261014", tournament_name="Testopolis",
                       tournament_date=pd.Timestamp("2026-10-14"), status="completed", has_box=True,
                       box_status="full", source="tml")]).to_parquet(s0 / "stage0_matches.parquet")
    pd.DataFrame([dict(match_id=mid, player_id="atp:1", player_name="atp:1", entry=None),
                  dict(match_id=mid, player_id="atp:91", player_name="atp:91", entry=None)]).to_parquet(
        s0 / "stage0_match_players.parquet")
    pd.DataFrame([dict(match_id=mid, lf=(pd.Timestamp("2026-10-16") - pd.Timestamp("1970-01-01")).days)]).to_parquet(
        tmp_path / "timing.parquet")
    pd.DataFrame([dict(tour="atp", tourney_id="2026-9001", round="R16", player_a="atp:1", player_b="atp:2",
                       day=day, tourney_name="Testopolis", start_utc="2026-10-20T10:00:00Z", stage0_id_a="atp:1",
                       stage0_id_b="atp:2") for day in ("2026-10-20", "garbage")]).to_csv(tmp_path / "m.csv",
                                                                                          index=False)
    out = DCK.data_check(tmp_path / "m.csv", stage0=s0, timing=tmp_path / "timing.parquet", held=set(),
                         packages=False, now=dt.datetime(2026, 10, 19, 12, 0, tzinfo=dt.timezone.utc), write=False,
                         show=False)
    assert out["Verdict"].tolist() == [DCK.FAILED, "ready"]                  # the failed check first
    assert out["Why"].iloc[0].startswith("ERROR while checking this match")
    assert "!!! ERROR while checking 1 match" in capsys.readouterr().out


# --- A5b, A7, A8 ---------------------------------------------------------------------------------------------------------
def test_a5b_an_error_repricing_one_match_reprices_the_others(monkeypatch, tmp_path, capsys):
    run = tmp_path / "live_2026-10-28_153009"
    run.mkdir()
    pd.DataFrame([{"match_id": f"atp|x|R16|atp:{i}|atp:{i + 10}", "tour": "atp", "player_a": f"Home Player{i}",
                   "player_b": f"Away Player{i}", "player_a_id": f"atp:{i}", "player_b_id": f"atp:{i + 10}"}
                  for i in (1, 2)]).to_csv(run / "predictions.csv", index=False)
    pd.DataFrame([{"player_a": f"Home Player{i}", "player_b": f"Away Player{i}", "pinnacle_matchup_id": mid,
                   "stage0_id_a": f"atp:{i}"} for i, mid in ((1, "not a number"), (2, 2))]).to_csv(
        run / "matchups_input.csv", index=False)
    monkeypatch.setattr(L, "latest_pull", lambda: _pull(["ATP Testville - R16", "ATP Testville - R16"]))
    out = pd.read_csv(L.live_prices(run, pull_now=False, path=tmp_path / "prices.csv"))
    assert out["match_id"].tolist() == ["atp|x|R16|atp:2|atp:12"]
    text = capsys.readouterr().out
    assert "!!! ERROR repricing 1 match: NOT repriced" in text and "Home Player1 v Away Player1 -- ValueError" in text


def test_a7_an_error_grading_one_run_grades_the_others(monkeypatch, tmp_path, capsys):
    runs = tmp_path / "runs"
    for name in ("live_2026-10-27_090000", "live_2026-10-28_090000"):
        (runs / name).mkdir(parents=True)
        pd.DataFrame({"match_id": [f"atp|{name}|R16|atp:1|atp:2"], "tour": "atp"}).to_csv(
            runs / name / "predictions.csv", index=False)
    s0 = tmp_path / "s0.parquet"
    pd.DataFrame({"match_id": [f"atp|{n}|R16|atp:1|atp:2" for n in ("live_2026-10-27_090000",
                                                                      "live_2026-10-28_090000")],
                  "status": "completed"}).to_parquet(s0)

    def grade(run, *, stage0=None, dry_run=False):
        if Path(run).name.endswith("27_090000"):
            raise R.StepFailed(R.StepResult("grade", "x", "failed", 1, 0.0, "logs/grade.log", ["boom"]))
        (Path(run) / "grade.json").write_text(json.dumps({"graded": 1, "mean_ll_model": 0.4, "mean_ll_market": 0.5}),
                                              encoding="utf-8")
    monkeypatch.setattr(PR, "grade", grade)
    out = tp.grade_due(stage0=s0, runs=runs).set_index("run")
    assert out.loc["live_2026-10-27_090000", "state"] == "ERROR"
    assert out.loc["live_2026-10-28_090000", ["state", "mean_ll_model"]].tolist() == ["graded now", 0.4]
    assert "!!! ERROR grading 1 run" in capsys.readouterr().out


def test_a8_a_record_that_cannot_be_read_is_said_and_the_others_read(tmp_path, capsys, monkeypatch):
    rec = tmp_path / "2026-10-28" / "records"
    rec.mkdir(parents=True)
    pd.DataFrame({"status": ["predicted"], "match_id": ["m1"], "run": ["live_x"]}).to_csv(
        rec / "bet_sheet_2026-10-28_1553.csv", index=False)
    (rec / "bet_sheet_2026-10-28_1600.csv").write_text('status,match_id\n"predicted,m2\n', encoding="utf-8")
    rows = CL.sheet_rows(tmp_path)
    assert rows["match_id"].tolist() == ["m1"]
    assert "!!! ERROR reading the bet sheet record" in capsys.readouterr().out
    monkeypatch.setattr(CL, "_row_metrics", lambda r, c, m: (_ for _ in ()).throw(KeyError("te_id")))
    assert CL._metrics_or_error({"player_a": "A", "player_b": "B"}, {}, {}) == {}
    assert "!!! ERROR measuring A v B" in capsys.readouterr().out
