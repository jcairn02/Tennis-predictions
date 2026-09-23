"""Shared helpers for the public data dump acquisition pipeline.

Layout (paths relative to the project root):

    data/raw/public_dump/<source>/          byte-exact downloads plus _receipts.jsonl (gitignored)
    data/dump/<branch>/<source>/...         tabular staging as Parquet, one file per source file
    data/studies/public_data_dump/          manifest, coverage summaries, download log (tracked)

Staging is organisation, not cleaning:

* every source column is kept, in source order, as a UTF-8 string; nothing is parsed or typed
* provenance columns start with an underscore: _source, _source_file, _row, _width_ok, _extra,
  _sha256, _retrieved_at, plus adapter-specific columns such as _route or _partition_value
* malformed rows are kept: short rows are padded with nulls, long rows keep their surplus
  fields joined with U+001F in _extra; _width_ok says which happened ("ok", "short", "long")
* source column names that start with an underscore are prefixed with "src" so they cannot
  collide with provenance columns; duplicate header names get a "__2", "__3" suffix
"""
from __future__ import annotations

import csv
import datetime as dt
import gzip
import hashlib
import io
import json
import re
import sys
import threading
import time
from pathlib import Path

import pyarrow as pa
import pyarrow.parquet as pq
import requests

PROJECT_ROOT = Path(__file__).resolve().parents[3]
RAW_ROOT = PROJECT_ROOT / "data" / "raw" / "public_dump"
DUMP_ROOT = PROJECT_ROOT / "data" / "dump"
STUDY_ROOT = PROJECT_ROOT / "data" / "studies" / "public_data_dump"
USER_AGENT = "tennis-predictions-public-dump/1.0 (research use; python-requests)"
DEFAULT_START_YEAR = 2010  # owner direction (22 Sep 2026): 2010+ matters, heavy emphasis on 2022-2026
DEFAULT_END_YEAR = 2026
EXTRA_SEP = "\x1f"
PROVENANCE_FIELDS = (
    ("_source", pa.string()),
    ("_source_file", pa.string()),
    ("_row", pa.int64()),
    ("_width_ok", pa.string()),
    ("_extra", pa.string()),
    ("_sha256", pa.string()),
    ("_retrieved_at", pa.string()),
)

csv.field_size_limit(min(sys.maxsize, 2**31 - 1))


# --------------------------------------------------------------------------------------
# small utilities
# --------------------------------------------------------------------------------------
def utc_now_iso() -> str:
    return dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds")


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def rel(path: Path) -> str:
    """Project-relative POSIX path when possible, else the absolute path."""
    try:
        return Path(path).resolve().relative_to(PROJECT_ROOT).as_posix()
    except ValueError:
        return Path(path).as_posix()


def slugify(value) -> str:
    """Filesystem-safe partition name from a verbatim label; never empty."""
    text = "" if value is None else str(value)
    text = re.sub(r"[^A-Za-z0-9]+", "_", text).strip("_").lower()
    return text or "unknown"


def years(start: int, end: int):
    if start > end:
        raise ValueError(f"start year {start} is after end year {end}")
    return range(start, end + 1)


def log(message: str) -> None:
    print(message, flush=True)


# --------------------------------------------------------------------------------------
# receipts and downloads
# --------------------------------------------------------------------------------------
class Receipts:
    """Append-only JSONL receipts for one source; the latest receipt per name wins."""

    def __init__(self, source: str, root: Path | None = None):
        self.source = source
        self.path = (root or RAW_ROOT) / source / "_receipts.jsonl"
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self._latest: dict[str, dict] = {}
        self._lock = threading.Lock()
        if self.path.exists():
            with self.path.open(encoding="utf-8") as handle:
                for line in handle:
                    if line.strip():
                        receipt = json.loads(line)
                        self._latest[receipt["name"]] = receipt

    def get(self, name: str) -> dict | None:
        return self._latest.get(name)

    def record(self, receipt: dict) -> dict:
        with self._lock:
            with self.path.open("a", encoding="utf-8") as handle:
                handle.write(json.dumps(receipt, ensure_ascii=False) + "\n")
            self._latest[receipt["name"]] = receipt
        return receipt

    def all(self) -> dict[str, dict]:
        return dict(self._latest)


_thread_local = threading.local()


def session() -> requests.Session:
    """One requests.Session per thread (Session objects are not thread-safe)."""
    sess = getattr(_thread_local, "session", None)
    if sess is None:
        sess = requests.Session()
        sess.headers["User-Agent"] = USER_AGENT
        _thread_local.session = sess
    return sess


def download(url: str, dest: Path, receipts: Receipts, *, name: str | None = None,
             refresh: bool = False, timeout: int = 120, retries: int = 2,
             validator=None, extra: dict | None = None, sess=None) -> dict:
    """GET ``url`` into ``dest`` and record a receipt. HTTP/validation failures never raise;
    the returned receipt carries an ``error`` key instead. A successful earlier download is
    reused (``skipped: True``) unless ``refresh`` is set."""
    name = name or rel(dest)
    previous = receipts.get(name)
    if dest.exists() and previous and not previous.get("error") and not refresh:
        return {**previous, "skipped": True}
    sess = sess or session()
    receipt = {"name": name, "url": url, "dest": rel(dest), "retrieved_at": utc_now_iso()}
    if extra:
        receipt.update(extra)
    last_error = None
    for attempt in range(retries + 1):
        try:
            response = sess.get(url, timeout=timeout)
            receipt.update(status=response.status_code, final_url=response.url,
                           content_type=response.headers.get("Content-Type"),
                           last_modified=response.headers.get("Last-Modified"),
                           etag=response.headers.get("ETag"))
            if response.status_code != 200:
                last_error = f"HTTP {response.status_code}"
                if response.status_code == 429 or 500 <= response.status_code < 600:
                    time.sleep(2 * (attempt + 1))
                    continue
                break
            body = response.content
            problem = validator(body) if validator else None
            if problem:
                last_error = problem
                break
            dest.parent.mkdir(parents=True, exist_ok=True)
            dest.write_bytes(body)
            receipt.update(bytes=len(body), sha256=sha256_bytes(body))
            return receipts.record(receipt)
        except requests.RequestException as err:
            last_error = f"{type(err).__name__}: {err}"
            time.sleep(2 * (attempt + 1))
    receipt["error"] = last_error or "unknown failure"
    return receipts.record(receipt)


def fetch_json(url: str, *, timeout: int = 60, sess=None):
    """GET a JSON document (no receipt; callers persist what they need)."""
    response = (sess or session()).get(url, timeout=timeout)
    response.raise_for_status()
    return response.json()


# validators ---------------------------------------------------------------------------
def not_html(body: bytes):
    head = body[:256].lstrip().lower()
    if head.startswith((b"<!doctype", b"<html", b"<head", b"<script")):
        return "HTML page returned instead of data"
    return None


def csv_header_contains(*needles: str):
    def check(body: bytes):
        problem = not_html(body)
        if problem:
            return problem
        first = body.split(b"\n", 1)[0].decode("utf-8", errors="replace").lstrip("﻿")
        missing = [n for n in needles if n not in first]
        return f"CSV header lacks {missing}: {first[:120]!r}" if missing else None
    return check


XLSX_MAGIC = b"PK\x03\x04"
XLS_MAGIC = b"\xd0\xcf\x11\xe0"


def workbook_magic(body: bytes):
    if body[:4] not in (XLSX_MAGIC, XLS_MAGIC):
        return f"not an Excel workbook (leading bytes {body[:4]!r})"
    if len(body) < 10_000:
        return f"workbook implausibly small ({len(body)} bytes)"
    return None


def gzip_magic(body: bytes):
    return None if body[:2] == b"\x1f\x8b" else f"not gzip (leading bytes {body[:2]!r})"


# --------------------------------------------------------------------------------------
# staging: rows -> Parquet with provenance
# --------------------------------------------------------------------------------------
class Provenance:
    """What every staged row records about where it came from."""

    def __init__(self, source: str, source_file: str, sha256: str | None, retrieved_at: str | None,
                 extra: dict | None = None):
        self.source = source
        self.source_file = source_file
        self.sha256 = sha256
        self.retrieved_at = retrieved_at
        self.extra = {k: (None if v is None else str(v)) for k, v in (extra or {}).items()}


def dedupe_header(header) -> list[str]:
    counts: dict[str, int] = {}
    out = []
    for index, raw in enumerate(header):
        name = ("" if raw is None else str(raw)).strip() or f"col_{index + 1}"
        if name.startswith("_"):
            name = "src" + name
        counts[name] = counts.get(name, 0) + 1
        out.append(name if counts[name] == 1 else f"{name}__{counts[name]}")
    return out


def _schema(columns: list[str], extra_columns) -> pa.Schema:
    fields = [pa.field(c, pa.string()) for c in columns]
    fields += [pa.field(n, t) for n, t in PROVENANCE_FIELDS]
    fields += [pa.field(k, pa.string()) for k in extra_columns]
    return pa.schema(fields)


def stage_rows(header, rows, dest: Path, prov: Provenance, batch_rows: int = 50_000) -> dict:
    """Write ``rows`` (an iterable of ``(row_number, fields)``) to ``dest`` as string Parquet.

    ``fields`` may be shorter or longer than the header; both are kept and flagged."""
    columns = dedupe_header(header)
    width = len(columns)
    schema = _schema(columns, list(prov.extra))
    dest.parent.mkdir(parents=True, exist_ok=True)
    stats = {"rows": 0, "short_rows": 0, "long_rows": 0}
    writer = pq.ParquetWriter(dest, schema, compression="zstd")
    batch: list[list] = [[] for _ in range(width)]
    b_row, b_width, b_extra = [], [], []

    def flush():
        if not b_row:
            return
        arrays = [pa.array(col, pa.string()) for col in batch]
        n = len(b_row)
        arrays += [
            pa.array([prov.source] * n, pa.string()),
            pa.array([prov.source_file] * n, pa.string()),
            pa.array(b_row, pa.int64()),
            pa.array(b_width, pa.string()),
            pa.array(b_extra, pa.string()),
            pa.array([prov.sha256] * n, pa.string()),
            pa.array([prov.retrieved_at] * n, pa.string()),
        ]
        arrays += [pa.array([value] * n, pa.string()) for value in prov.extra.values()]
        writer.write_table(pa.Table.from_arrays(arrays, schema=schema))
        for col in batch:
            col.clear()
        b_row.clear()
        b_width.clear()
        b_extra.clear()

    try:
        for row_number, fields in rows:
            fields = list(fields)
            if len(fields) < width:
                fields += [None] * (width - len(fields))
                flag, extra = "short", None
                stats["short_rows"] += 1
            elif len(fields) > width:
                extra = EXTRA_SEP.join("" if f is None else str(f) for f in fields[width:])
                fields = fields[:width]
                flag = "long"
                stats["long_rows"] += 1
            else:
                flag, extra = "ok", None
            for col, value in zip(batch, fields):
                col.append(None if value is None else str(value))
            b_row.append(row_number)
            b_width.append(flag)
            b_extra.append(extra)
            stats["rows"] += 1
            if len(b_row) >= batch_rows:
                flush()
        flush()
    finally:
        writer.close()
    stats.update(columns=columns, dest=rel(dest), bytes=dest.stat().st_size)
    return stats


def open_text(path: Path) -> tuple[str, str]:
    """Decode a (possibly gzipped) text file; returns (text, encoding_used)."""
    raw = path.read_bytes()
    if path.suffix == ".gz":
        raw = gzip.decompress(raw)
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1", errors="replace"), "latin-1-replace"


def iter_delimited(path: Path, delimiter: str = ","):
    """Yield (header, encoding, row iterator of (row_number, fields)) for a CSV file.

    Row numbers count data records as parsed by csv.reader (1-based), not physical lines;
    completely empty lines are skipped and counted separately by the caller if needed."""
    text, encoding = open_text(path)
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    header = next(reader, [])

    def rows():
        number = 0
        for fields in reader:
            if not fields:
                continue
            number += 1
            yield number, fields

    return header, encoding, rows()


def stage_delimited(src: Path, dest: Path, prov: Provenance, delimiter: str = ",") -> dict:
    header, encoding, rows = iter_delimited(src, delimiter)
    stats = stage_rows(header, rows, dest, prov)
    stats["encoding"] = encoding
    stats["source_path"] = rel(src)
    return stats


# Excel -------------------------------------------------------------------------------
def cell_to_str(value):
    """Verbatim-as-possible string form of an Excel cell value."""
    if value is None:
        return None
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, dt.datetime):
        if (value.hour, value.minute, value.second, value.microsecond) == (0, 0, 0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float):
        if value.is_integer() and abs(value) < 1e15:
            return str(int(value))
        return repr(value)
    return str(value)


def read_workbook(path: Path) -> tuple[str, list, list[list], dict]:
    """First worksheet of an .xlsx/.xls workbook as (sheet_name, header, rows, meta).

    Cells are converted with cell_to_str; the header is the first row with any content."""
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        import openpyxl

        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheet = book.worksheets[0]
        raw_rows = [list(r) for r in sheet.iter_rows(values_only=True)]
        sheet_names = book.sheetnames
        book.close()
    elif suffix == ".xls":
        import xlrd

        book = xlrd.open_workbook(path)
        sheet = book.sheet_by_index(0)
        sheet_names = book.sheet_names()
        raw_rows = []
        for r in range(sheet.nrows):
            cells = []
            for cell in sheet.row(r):
                if cell.ctype == xlrd.XL_CELL_EMPTY or cell.ctype == xlrd.XL_CELL_BLANK:
                    cells.append(None)
                elif cell.ctype == xlrd.XL_CELL_DATE:
                    cells.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    cells.append(bool(cell.value))
                elif cell.ctype == xlrd.XL_CELL_ERROR:
                    cells.append(f"#ERR{cell.value}")
                else:
                    cells.append(cell.value)
            raw_rows.append(cells)
    else:
        raise ValueError(f"unsupported workbook type: {path}")
    header_index = next((i for i, r in enumerate(raw_rows) if any(c not in (None, "") for c in r)), None)
    if header_index is None:
        return sheet_names[0], [], [], {"sheet_names": sheet_names, "blank_rows": len(raw_rows)}
    header = raw_rows[header_index]
    while header and header[-1] in (None, ""):
        header.pop()
    rows = []
    blank = 0
    for r in raw_rows[header_index + 1:]:
        values = [cell_to_str(v) for v in r]
        while values and values[-1] is None:
            values.pop()
        if not values:
            blank += 1
            continue
        rows.append(values)
    meta = {"sheet_names": sheet_names, "sheet": sheet_names[0], "header_row_index": header_index,
            "blank_rows": blank}
    return sheet_names[0], [cell_to_str(h) for h in header], rows, meta


# Parquet passthrough ------------------------------------------------------------------
def stage_parquet_as_strings(src: Path, dest: Path, prov: Provenance) -> dict:
    """Copy a Parquet table into the staging layout with every column rendered as a string.
    Nested values (lists/structs) become JSON text; dates/timestamps become ISO text."""
    table = pq.read_table(src)
    header = table.column_names
    converted = []
    for column in table.columns:
        kind = column.type
        if pa.types.is_list(kind) or pa.types.is_large_list(kind) or pa.types.is_struct(kind) or pa.types.is_map(kind):
            converted.append([None if v is None else json.dumps(v, ensure_ascii=False, default=str)
                              for v in column.to_pylist()])
        elif pa.types.is_date(kind) or pa.types.is_timestamp(kind):
            converted.append([None if v is None else v.isoformat() for v in column.to_pylist()])
        elif pa.types.is_boolean(kind):
            converted.append([None if v is None else ("TRUE" if v else "FALSE") for v in column.to_pylist()])
        else:
            converted.append(column.cast(pa.string()).to_pylist())
    rows = ((i + 1, [col[i] for col in converted]) for i in range(table.num_rows))
    stats = stage_rows(header, rows, dest, prov)
    stats["source_path"] = rel(src)
    stats["source_schema"] = {name: str(t) for name, t in zip(table.column_names, table.schema.types)}
    return stats


# metadata sidecars ----------------------------------------------------------------------
def write_meta(dest: Path, meta: dict) -> Path:
    path = dest.parent / (dest.name + ".meta.json")
    path.write_text(json.dumps(meta, indent=1, ensure_ascii=False), encoding="utf-8")
    return path


def staged_meta(branch: str, source: str, dest: Path, receipt: dict, stats: dict, **extra) -> dict:
    meta = {"branch": branch, "source": source, "dest": rel(dest),
            "rows": stats.get("rows"), "columns": stats.get("columns"), "bytes": stats.get("bytes"),
            "short_rows": stats.get("short_rows"), "long_rows": stats.get("long_rows"),
            "encoding": stats.get("encoding"), "source_path": stats.get("source_path"),
            "url": receipt.get("url"), "final_url": receipt.get("final_url"),
            "sha256": receipt.get("sha256"), "source_bytes": receipt.get("bytes"),
            "retrieved_at": receipt.get("retrieved_at"), "last_modified": receipt.get("last_modified")}
    meta.update(extra)
    write_meta(dest, meta)
    return meta


def stage_csv_file(branch: str, source: str, receipt: dict, src: Path, dest: Path,
                   extra_provenance: dict | None = None, **meta_extra) -> dict:
    """Stage one downloaded CSV (optionally gzipped) and write its sidecar."""
    prov = Provenance(source, rel(src), receipt.get("sha256"), receipt.get("retrieved_at"), extra_provenance)
    stats = stage_delimited(src, dest, prov)
    return staged_meta(branch, source, dest, receipt, stats, **meta_extra)


def dump_path(branch: str, source: str, *parts: str) -> Path:
    return DUMP_ROOT.joinpath(branch, source, *parts)


def raw_path(source: str, *parts: str) -> Path:
    return RAW_ROOT.joinpath(source, *parts)


def summarize(receipts: Receipts, staged: list[dict], skipped: list[dict] | None = None) -> dict:
    all_receipts = receipts.all().values()
    errors = [r for r in all_receipts if r.get("error")]
    return {"source": receipts.source, "downloads": len(all_receipts), "errors": errors,
            "staged_files": len(staged), "staged_rows": sum(s.get("rows") or 0 for s in staged),
            "skipped": skipped or []}
