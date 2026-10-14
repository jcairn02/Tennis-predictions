"""The bet sheet: section A's output, organised for the reader deciding what to bet (the owner's request of 13 Oct 2026;
docs/studies/master_pipeline/process_log.md, Part 2).

bet_diagnostics() (diagnostics.py) computes the record: one row per match, every figure on Stage 0's sides A and B. This
module only reads the record and says it the way a bettor reads it. It computes no probability, EV or block.

    the sheet       one row per predicted match: BET (the player the model backs at Pinnacle's odds, with the stake),
                    then PASS (no side is +EV at Pinnacle's odds: the odds that would make the model's side a bet), then
                    BLOCKED (a safety rule fired); each with the model's and the market's chance, the edge, the expected
                    profit, the start in local time and a plain note
    the header      the totals, the price's time, the model, the stake rule, and what applies to the whole run, once

It is shown in the notebook (show) and written as a workbook (write_workbook): "Bet sheet", then "Matches" (every figure
of both players, in Pinnacle's order), "Not predicted" (the listed matches left out, and why), "Run" (provenance) and
"How to read" (every column).
"""
from __future__ import annotations

import datetime as dt
import html
import math
import re
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np
import pandas as pd

THIN_EV = 0.01            # a bet whose expected profit is under 1% is flagged "thin" (a display flag, not a betting rule)
ORDER = {"BET": 0, "PASS": 1, "BLOCKED": 2}
STAKE_COLUMNS = {"owner_8oct": "stake_units_owner_8oct_research"}     # the record's stake column of each rule
STAKE_RULES = {"owner_8oct": "your 8 Oct research rule, in units: favourites are staked to win 1 unit, 1 unit at even "
                             "money, underdogs d / (2(d - 1)) units (0.83 at 2.50, 0.75 at 3.00). It sizes by the price "
                             "only, not by the edge. Not a decided staking rule"}
MODELS = {"tabpfn35": "the current best model (TabPFN-3.5 on 57 inputs, then stage 2: calibration and serve-speed form, "
                      "plus the five-pattern adjustment when it is switched on)",
          "tabpfn35_thinking": "TabPFN-3.5-Thinking on 57 inputs, then stage 2"}

COLOURS = {"BET": ("#E2EFDA", "#1E4620"), "PASS": ("#F2F2F2", "#3A3A3A"), "BLOCKED": ("#FBE5D6", "#7A2E0E")}
HEAD = ("#1F3864", "#FFFFFF")
GROUP = ("#D9E1F2", "#1F3864")

# How each kind of column is written: Excel's number format, and the notebook's
FORMATS = {"pct": ("0.0%", "{:.1%}"), "pts": ("+0.0;-0.0;0.0", "{:+.1f}"), "ev": ("+0.0%;-0.0%;0.0%", "{:+.1%}"),
           "odds": ("0.000", "{:.3f}"), "us": ("+0;-0;0", "{:+.0f}"), "units": ("0.00", "{:.2f}"),
           "int": ("0", "{:.0f}"), "money": ("#,##0", "{:,.0f}"), "kmh": ("+0.0;-0.0;0.0", "{:+.1f}"),
           "spread": ("0.0", "{:.1f}"), "text": (None, None)}
WRAP = {"Notes": 72}       # columns written wrapped, at this width (Excel's characters)


# ---------------------------------------------------------------------------------------------------------------------
# small readers of a record row

def _f(x) -> float:
    try:
        return float(x)
    except (TypeError, ValueError):
        return float("nan")


def _text(x) -> str:
    return x.strip() if isinstance(x, str) else ""


def _bool(x):
    if isinstance(x, (bool, np.bool_)):
        return bool(x)
    if isinstance(x, str) and x.strip().lower() in ("true", "false"):
        return x.strip().lower() == "true"
    return None


def _sentence(s: str) -> str:
    s = s.strip()
    return (s[0].upper() + s[1:] + ("" if s.endswith(".") else ".")) if s else ""


def american(d) -> float:
    """Decimal odds as American odds (+132, -318); NaN when not a price."""
    d = _f(d)
    if not np.isfinite(d) or d <= 1:
        return float("nan")
    return float(round((d - 1) * 100)) if d >= 2 else float(round(-100 / (d - 1)))


def local_time(ts) -> dt.datetime | None:
    """A timestamp in this machine's local time (its own offset, so a DST change between now and the match is right);
    a value without a zone is read as UTC."""
    if ts is None or (isinstance(ts, float) and math.isnan(ts)) or (isinstance(ts, str) and not ts.strip()):
        return None
    t = pd.Timestamp(ts)
    if pd.isna(t):
        return None
    if t.tzinfo is None:
        t = t.tz_localize("UTC")
    return t.to_pydatetime().astimezone()


def zone_label(t: dt.datetime | None = None) -> str:
    """The local zone's short name ("PDT" from Windows' "Pacific Daylight Time")."""
    t = t if t is not None and t.tzinfo else dt.datetime.now().astimezone()
    name = t.tzname() or "local"
    return "".join(w[0] for w in name.split()).upper() if " " in name else name


def when(ts, *, year: bool = False) -> str:
    """Sun 11 Oct 21:00, local time."""
    t = local_time(ts)
    if t is None:
        return ""
    return f"{t:%a} {t.day} {t:%b}" + (f" {t.year}" if year else "") + f" {t:%H:%M}"


def decision(r) -> str:
    """BET (a pick), PASS (predicted, not blocked, no side +EV at Pinnacle's odds), BLOCKED, or NOT PREDICTED."""
    if r.get("status") != "predicted":
        return "NOT PREDICTED"
    if _bool(r.get("blocked")):
        return "BLOCKED"
    return "BET" if r.get("pick") in ("A", "B") else "PASS"


def lean(r) -> str | None:
    """The side the model rates above the market: the only side that can be +EV at a bookmaker's odds."""
    e = _f(r.get("edge_a"))
    if not np.isfinite(e) or e == 0:
        return None
    return "A" if e > 0 else "B"


def side(r, s: str) -> dict:
    """One side's figures, from the record's side-A/B columns."""
    a = s == "A"
    pm, base = _f(r.get("p_market_a")), _f(r.get("p_base_a"))
    off = _f(r.get("q_a_patterns_off"))
    return {"name": _text(r.get("player_a" if a else "player_b")),
            "id": _text(r.get("player_a_id" if a else "player_b_id")),
            "odds": _f(r.get("pinnacle_odds_a" if a else "pinnacle_odds_b")),
            "market": pm if a else 1 - pm,
            "model": _f(r.get("q_a" if a else "q_b")),
            "model_off": off if a else 1 - off,
            "base": base if a else 1 - base,
            "ev": _f(r.get("ev_a_pinnacle" if a else "ev_b_pinnacle")),
            "fair": _f(r.get("fair_odds_a" if a else "fair_odds_b")),
            "readings": _f(r.get("speed_readings_28d_a" if a else "speed_readings_28d_b"))}


def listed(r) -> tuple[str, str]:
    """The two sides in Pinnacle's order (its home player first); Stage 0's order when unknown."""
    return ("B", "A") if r.get("home_side") == "B" else ("A", "B")


def match_name(r) -> str:
    one, two = listed(r)
    return f"{side(r, one)['name']} v {side(r, two)['name']}"


def _court(r) -> str:
    return ", ".join(x for x in (_text(r.get("surface")),
                                 {True: "indoor", False: "outdoor"}.get(_bool(r.get("indoor")), "")) if x)


def event(r) -> str:
    """ATP Shanghai R32 · Hard, outdoor · best of 3."""
    tour, name = _text(r.get("tour")).upper(), _text(r.get("tournament"))
    if tour and not name.upper().startswith(tour):
        name = f"{tour} {name}"
    parts = [f"{name} {_text(r.get('round'))}".strip()]
    if _court(r):
        parts.append(_court(r))
    bo = _f(r.get("best_of"))
    if np.isfinite(bo):
        parts.append(f"best of {int(bo)}")
    return " · ".join(p for p in parts if p)


# ---------------------------------------------------------------------------------------------------------------------
# the record's reasons, said plainly (an unknown text is kept as it is)

_BLOCKS = [
    (r"^gate: do not bet: inputs incomplete", None),
    (r"^quote (\d+) min old \(over (\d+)\)", "The price is {0} min old (limit {1}): reprice with cell A6 before betting."),
    (r"^the match has started", "The match has started."),
    (r"^not a main-tour match", "Not a main-tour match: outside what the model was trained on."),
    (r"^unsupported match format", "A match format the model does not cover."),
    (r"^price read-back mismatch", "Safety check failed: the price the model used is not the listed price. Do not bet; "
                                   "run A3 to A5 again."),
    (r"^players read back differ", "Safety check failed: the players the model priced are not the listed players. Do "
                                   "not bet."),
    (r"^new price off the curve", "The new price is outside the range this run priced (5% to 95%): run A4 again."),
    (r"^no bookmaker odds", "No bookmaker odds with their margin, so no expected profit can be worked out."),
]

_LEFT_OUT = [
    (r"^no main-tour Tennis My Life event named like", "First round of a new event: the model needs one played match "
                                                       "at the event first."),
    (r"matches several events", "The event's name matches several events: left out to be safe."),
    (r"final is played", "Looks like a new event of the same name (the old one's final is played)."),
    (r"latest match is .* over a week before", "The event's results are over a week old."),
    (r"^already played", "Already played."),
    (r"^round check", "The round could not be confirmed from the event's results."),
    (r"^'(.+)' names (\d+) Stage 0 players", "{0}'s name matches {1} players in the data: cannot be predicted until "
                                             "the duplicate is merged."),
    (r"^'(.+)' is unknown to Stage 0", "{0} is not in the data yet (an alias may be needed)."),
    (r"^round .* not understood", "Pinnacle's round label was not understood."),
    (r"level \(outside the model", "Not a main-tour event (Challenger, 125 or ITF level)."),
    (r"\(ITF or other\)", "Not a main-tour event (ITF or other)."),
    (r"^qualifying$", "Qualifying."),
    (r"team event", "A team event."),
    (r"^doubles$", "Doubles."),
    (r"in-play", "Already in play."),
    (r"two listed matches of the same pair", "Listed twice by Pinnacle: ambiguous."),
    (r"build_matchups only=", "Not chosen (only= in cell A3)."),
    (r"build_matchups drop=", "Dropped (drop= in cell A3)."),
]


def _parts(text) -> list[str]:
    return [p.strip() for p in _text(text).split("; ") if p.strip()]


def why_blocked(r) -> str:
    out = []
    for part in _parts(r.get("block_reasons")):
        for pat, plain in _BLOCKS:
            m = re.match(pat, part)
            if m:
                if plain is None:
                    gaps = _text(r.get("gate_detail"))
                    out.append(f"Data gap: no usable box score for {gaps}, so the model's inputs are incomplete." if gaps
                               else "Data gap: an earlier match of a player has no usable box score, so the model's "
                                    "inputs are incomplete.")
                else:
                    out.append(plain.format(*m.groups()))
                break
        else:
            out.append(_sentence(part))
    return " ".join(out)


def why_left_out(reason) -> str:
    text = _text(reason)
    for pat, plain in _LEFT_OUT:
        m = re.search(pat, text)
        if m:
            return plain.format(*m.groups())
    return _sentence(text)


def match_warnings(r) -> list[str]:
    """The record's match-specific warnings, plainly (the run's own notes are said once, in the header)."""
    out = []
    for part in _parts(r.get("warnings")):
        m = re.match(r"^the 16 seeds disagree \(sd ([0-9.]+)\)", part)
        if m:
            out.append(f"The model's 16 runs disagree (spread ±{100 * float(m.group(1)):.1f} pts): less certain than "
                       "usual.")
        elif part.startswith("serve-speed form missing"):
            names = [side(r, s)["name"] for s in ("A", "B") if side(r, s)["readings"] == 0]
            out.append(f"No serve-speed readings in 28 days for {' and '.join(names) or 'a player'}: the serve-speed "
                       "step has no signal there.")
        elif part.startswith("price outside the curve"):
            out.append("Long odds: a later price cannot be read off this run (a reprice needs a new run).")
        elif part.startswith("the quote time carries no time zone"):
            out.append("The quote time had no time zone (read as UTC).")
        else:
            out.append(_sentence(part))
    return out


def no_edge_at(r) -> str:
    """Where the model's edge on its side would vanish: the market chance (and fair odds) nearest the current price at
    which the model's curve crosses the price."""
    s = lean(r)
    if s is None:
        return ""
    near = _f(r.get("nearest_break_even_p_a"))
    if np.isfinite(near):
        p = near if s == "A" else 1 - near
        n = len(_text(r.get("break_even_p_a")).split(";"))
        return f"{p:.1%} (odds {1 / p:.3f})" + (f"; {n - 1} more crossing{'s' * (n > 2)}" if n > 1 else "")
    note = _text(r.get("break_even_note"))
    if note.startswith("no crossing"):
        return f"never between 5% and 95%: the model prefers {side(r, s)['name']} at every price"
    return note


# ---------------------------------------------------------------------------------------------------------------------
# the sheet

@dataclass
class Sheet:
    made_at: dt.datetime
    zone: str
    title: str
    header: list[str]
    notes: list[str]
    table: pd.DataFrame
    kinds: dict[str, str]
    matches: pd.DataFrame
    matches_kinds: dict[str, str]
    groups: list[tuple[str, int]]
    left_out: pd.DataFrame
    left_out_summary: str
    run: list[tuple[str, str]]
    files: dict = field(default_factory=dict)


def _start_key(r) -> float:
    t = local_time(r.get("start_utc"))
    return t.timestamp() if t is not None else math.inf


def _event_name(r) -> str:
    """ATP Shanghai (a left-out row's Pinnacle league "WTA Wuhan - R1" without its round)."""
    tour, name = _text(r.get("tour")).upper(), re.sub(r"\s+-\s+.*$", "", _text(r.get("tournament")))
    return name if (not tour or name.upper().startswith(tour)) else f"{tour} {name}"


def _venue_columns(record: pd.DataFrame, venue: str) -> tuple[str, str] | None:
    key = venue or "venue"
    return (f"{(venue or 'Venue').capitalize()} payout", f"{(venue or 'Venue').capitalize()} expected profit") \
        if f"ev_a_{key}" in record.columns or f"ev_b_{key}" in record.columns else None


def notes_for(r, venue: str = "") -> str:
    """One plain note per match: why it is a PASS or BLOCKED, what weakens a bet, the match's own warnings."""
    d = decision(r)
    out = []
    eff = _text(r.get("patterns_effect"))
    if eff == "bet only with the patterns":
        out.append("Only the pattern adjustment makes this a bet: without it the model would pass. Keep the stake small.")
    elif eff == "the patterns remove this bet":
        out.append("Without the pattern adjustment this would be a bet; with it, the model passes.")
    elif eff == "the patterns switch the side":
        out.append("The pattern adjustment switches the player the model backs.")
    s = r.get("pick") if d == "BET" else (lean(r) if d == "PASS" else None)
    if d == "BET":
        ev = side(r, s)["ev"]
        if np.isfinite(ev) and ev < THIN_EV:
            out.append(f"Thin: expected profit under {THIN_EV:.0%}; a small price move removes it.")
    elif d == "PASS" and s in ("A", "B"):
        v = side(r, s)
        out.append(f"No value at {v['odds']:.3f}: the model's edge on {v['name']} "
                   f"({100 * (v['model'] - v['market']):+.1f} pts) is smaller than the bookmaker's margin, so the "
                   f"expected profit is {v['ev']:+.1%}. Worth a bet at {v['fair']:.3f} or more.")
        vev = _f(r.get(f"ev_{s.lower()}_{venue or 'venue'}"))
        if np.isfinite(vev) and vev > 0:
            out.append(f"At {venue or 'the other venue'} it is +EV ({vev:+.1%}).")
    elif d == "PASS":
        out.append("The model agrees with the market: no side has an edge.")
    elif d == "BLOCKED":
        out.append(why_blocked(r))
    return " ".join(out + match_warnings(r))


def _sheet_table(rec: pd.DataFrame, zone: str, stake_rule: str | None, venue: str) -> tuple[pd.DataFrame, dict]:
    start = f"Starts ({zone})"
    vcols = _venue_columns(rec, venue)
    key = venue or "venue"
    stake_col = STAKE_COLUMNS.get(stake_rule or "")
    rows = []
    for _, r in rec.iterrows():
        d = decision(r)
        if d == "NOT PREDICTED":
            continue
        row = {start: when(r.get("start_utc")), "Match": match_name(r), "Decision": d, "_order": ORDER[d],
               "_start": _start_key(r)}
        s = r.get("pick") if d == "BET" else (lean(r) if d == "PASS" else None)
        if s in ("A", "B"):
            v = side(r, s)
            stake = _f(r.get(stake_col)) if (d == "BET" and stake_col) else float("nan")
            row.update({"Player": v["name"], "Stake (units)": stake, "Odds": v["odds"], "US odds": american(v["odds"]),
                        "Lowest odds to take": v["fair"], "Model chance": v["model"],
                        "Model chance without patterns": v["model_off"], "Market chance": v["market"],
                        "Edge (pts)": 100 * (v["model"] - v["market"]), "Expected profit": v["ev"]})
            if vcols:
                row[vcols[0]] = _f(r.get(f"venue_payout_{s.lower()}"))
                row[vcols[1]] = _f(r.get(f"ev_{s.lower()}_{key}"))
        row.update({"Event": event(r), "Notes": notes_for(r, venue)})
        rows.append(row)
    cols = [start, "Match", "Decision", "Player", "Stake (units)", "Odds", "US odds", "Lowest odds to take",
            "Model chance", "Model chance without patterns", "Market chance", "Edge (pts)", "Expected profit",
            *(vcols or ()), "Event", "Notes"]
    t = pd.DataFrame(rows, columns=[*cols, "_order", "_start"])
    if len(t):
        t = t.sort_values(["_order", "_start", "Match"], kind="stable")
    t = t[cols].reset_index(drop=True)
    if not stake_col:
        t = t.drop(columns=["Stake (units)"])
    kinds = {"Stake (units)": "units", "Odds": "odds", "US odds": "us", "Lowest odds to take": "odds",
             "Model chance": "pct", "Model chance without patterns": "pct", "Market chance": "pct", "Edge (pts)": "pts",
             "Expected profit": "ev"}
    if vcols:
        kinds.update({vcols[0]: "odds", vcols[1]: "ev"})
    return t, {c: kinds.get(c, "text") for c in t.columns}


def _matches_table(rec: pd.DataFrame, zone: str, stake_rule: str | None, speed_km_h: float
                   ) -> tuple[pd.DataFrame, list[tuple[str, int]], dict]:
    """Every figure of each predicted match, both players in Pinnacle's order (1 = listed first)."""
    stake_col = STAKE_COLUMNS.get(stake_rule or "")
    groups = [
        ("The match", [(f"Starts ({zone})", "text"), ("Player 1", "text"), ("Player 2", "text"), ("Tour", "text"),
                       ("Event", "text"), ("Round", "text"), ("Court", "text"), ("Best of", "int"),
                       ("Starts (UTC)", "text")]),
        ("The decision", [("Decision", "text"), ("Bet on", "text"), ("Pattern effect", "text"),
                          ("Stake (units)", "units"), ("Lowest odds to take", "odds"), ("Data check", "text"),
                          ("Notes", "text")]),
        ("Pinnacle's price", [("Odds 1", "odds"), ("Odds 2", "odds"), ("US odds 1", "us"), ("US odds 2", "us"),
                              ("Margin", "pct"), ("Market chance 1", "pct"), ("Market chance 2", "pct"),
                              (f"Quoted ({zone})", "text"), ("Price age (min)", "int"), ("Limit (max stake)", "money")]),
        ("The model", [("Model chance 1", "pct"), ("Model chance 2", "pct"),
                       ("Model chance 1 without patterns", "pct"), ("Pattern shift 1 (pts)", "pts"),
                       ("Fair odds 1", "odds"), ("Fair odds 2", "odds"), ("TabPFN alone 1", "pct"),
                       ("Stage 2 shift 1 (pts)", "pts"),
                       ("Serve-speed gap 1-2 (km/h)", "kmh"), ("Speed readings 1 (28 d)", "int"),
                       ("Speed readings 2 (28 d)", "int"), ("Seed spread (± pts)", "spread")]),
        ("The value", [("Model's side", "text"), ("Edge (pts)", "pts"), ("Expected profit 1", "ev"),
                       ("Expected profit 2", "ev"), ("Edge gone if the market reaches", "text")]),
        ("Identifiers", [("Match id", "text"), ("Player 1 id", "text"), ("Player 2 id", "text"), ("Run", "text")]),
    ]
    rows = []
    for _, r in rec.iterrows():
        d = decision(r)
        if d == "NOT PREDICTED":
            continue
        one, two = listed(r)
        p1, p2 = side(r, one), side(r, two)
        s = lean(r)
        spd = _f(r.get("spd_L")) * (1 if one == "A" else -1) * speed_km_h
        sd = _f(r.get("seed_sd"))
        bet = r.get("pick") if d == "BET" else None
        gate = _text(r.get("gate"))
        qt = r.get("quoted_at")
        rows.append({
            f"Starts ({zone})": when(r.get("start_utc")), "Starts (UTC)": _text(r.get("start_utc")),
            "Tour": _text(r.get("tour")).upper(), "Event": _text(r.get("tournament")), "Round": _text(r.get("round")),
            "Court": _court(r), "Best of": _f(r.get("best_of")), "Player 1": p1["name"], "Player 2": p2["name"],
            "Decision": d, "Bet on": side(r, bet)["name"] if bet in ("A", "B") else "",
            "Pattern effect": _text(r.get("patterns_effect")),
            "Model chance 1 without patterns": p1["model_off"],
            "Pattern shift 1 (pts)": 100 * (p1["model"] - p1["model_off"]) if np.isfinite(p1["model_off"])
            else float("nan"),
            "Stake (units)": _f(r.get(stake_col)) if (bet and stake_col) else float("nan"),
            "Lowest odds to take": side(r, s)["fair"] if (s and d != "BLOCKED") else float("nan"),
            "Data check": "complete" if gate == "may be bet" else ("incomplete" if "incomplete" in gate else gate),
            "Notes": notes_for(r),
            "Odds 1": p1["odds"], "Odds 2": p2["odds"], "US odds 1": american(p1["odds"]),
            "US odds 2": american(p2["odds"]), "Margin": _f(r.get("overround")) - 1,
            "Market chance 1": p1["market"], "Market chance 2": p2["market"], f"Quoted ({zone})": when(qt),
            "Price age (min)": _f(r.get("quote_age_min")), "Limit (max stake)": _f(r.get("max_risk_stake")),
            "Model chance 1": p1["model"], "Model chance 2": p2["model"], "Fair odds 1": p1["fair"],
            "Fair odds 2": p2["fair"], "TabPFN alone 1": p1["base"],
            "Stage 2 shift 1 (pts)": 100 * (p1["model"] - p1["base"]), "Serve-speed gap 1-2 (km/h)": spd,
            "Speed readings 1 (28 d)": p1["readings"], "Speed readings 2 (28 d)": p2["readings"],
            "Seed spread (± pts)": 100 * sd if np.isfinite(sd) else float("nan"),
            "Model's side": side(r, s)["name"] if s else "",
            "Edge (pts)": 100 * (side(r, s)["model"] - side(r, s)["market"]) if s else 0.0,
            "Expected profit 1": p1["ev"], "Expected profit 2": p2["ev"],
            "Edge gone if the market reaches": no_edge_at(r),
            "Match id": _text(r.get("match_id")), "Player 1 id": p1["id"], "Player 2 id": p2["id"],
            "Run": _text(r.get("run")), "_order": ORDER[d], "_start": _start_key(r)})
    cols = [c for _, cs in groups for c, _ in cs]
    t = pd.DataFrame(rows, columns=[*cols, "_order", "_start"])
    if len(t):
        t = t.sort_values(["_order", "_start", "Player 1"], kind="stable")
    t = t[cols].reset_index(drop=True)
    kinds = {c: k for _, cs in groups for c, k in cs}
    if not stake_col:
        t = t.drop(columns=["Stake (units)"])
        groups = [(g, [(c, k) for c, k in cs if c != "Stake (units)"]) for g, cs in groups]
        kinds.pop("Stake (units)")
    return t, [(g, len(cs)) for g, cs in groups], kinds


def _left_out_table(rec: pd.DataFrame, zone: str) -> tuple[pd.DataFrame, str]:
    out = rec[rec["status"].eq("not predicted")] if "status" in rec else rec.iloc[0:0]
    rows = []
    for _, r in out.iterrows():
        rows.append({f"Starts ({zone})": when(r.get("start_utc")), "Tour": _text(r.get("tour")).upper(),
                     "Event": event(r), "Match": match_name(r), "Odds 1": _f(r.get("pinnacle_odds_a")),
                     "Odds 2": _f(r.get("pinnacle_odds_b")), "Why not predicted": why_left_out(r.get("block_reasons")),
                     "Detail": _text(r.get("block_reasons")), "_start": _start_key(r)})
    t = pd.DataFrame(rows, columns=[f"Starts ({zone})", "Tour", "Event", "Match", "Odds 1", "Odds 2",
                                    "Why not predicted", "Detail", "_start"])
    if len(t):
        t = t.sort_values(["_start", "Event"], kind="stable")
    names = [_event_name(r) for _, r in out.iterrows()]
    t = t.drop(columns="_start").reset_index(drop=True)
    if t.empty:
        return t, ""
    by = pd.DataFrame({"why": [why_left_out(x) for x in out["block_reasons"]], "event": names})
    parts = [f"{len(g)} × {why.rstrip('.')} ({', '.join(sorted(set(g['event'])))})"
             for why, g in sorted(by.groupby("why", sort=False), key=lambda kv: -len(kv[1]))]
    return t, f"{len(t)} listed main-tour matches were not predicted: " + "; ".join(parts) + "."


def build(record: pd.DataFrame, *, made_at: dt.datetime, stake_rule: str | None = None, venue: str = "",
          notes: list[dict] | None = None, meta: dict | None = None) -> Sheet:
    """The sheet of a record (bet_diagnostics' rows). notes: the run's notes ({"plain": ..., "detail": ...}), said
    once. meta: run, model, instance, checkpoint, seeds, stale_minutes, reprice_file, history_line, speed_km_h."""
    meta, notes = meta or {}, notes or []
    made = made_at.astimezone() if made_at.tzinfo else made_at.replace(tzinfo=dt.timezone.utc).astimezone()
    zone = zone_label(made)
    table, kinds = _sheet_table(record, zone, stake_rule, venue)
    matches, groups, mkinds = _matches_table(record, zone, stake_rule, _f(meta.get("speed_km_h")))
    left, left_summary = _left_out_table(record, zone)
    pred = record[record["status"].eq("predicted")] if "status" in record else record.iloc[0:0]

    n = table["Decision"].value_counts().to_dict() if len(table) else {}
    nb, npass, nblk = n.get("BET", 0), n.get("PASS", 0), n.get("BLOCKED", 0)
    events = []
    for _, r in pred.iterrows():
        e = f"{_text(r.get('tour')).upper()} {_text(r.get('tournament'))}".strip()
        if e not in events:
            events.append(e)
    title = f"Tennis bet sheet · {', '.join(events) or 'no predicted match'} · made {when(made, year=True)} {zone}"
    starts = [t for t in (local_time(x) for x in pred.get("start_utc", pd.Series(dtype=object))) if t is not None]
    header = []
    counts = f"{nb} to bet, {npass} pass, {nblk} blocked"
    if len(left):
        counts += f"; {len(left)} listed matches not predicted (sheet 'Not predicted')"
    header.append(counts + (f". Matches start {when(min(starts))} to {when(max(starts))} {zone}." if starts else "."))
    bets = table[table["Decision"].eq("BET")] if len(table) else table
    if nb and "Stake (units)" in table:
        st = bets["Stake (units)"]
        win = (st * (bets["Odds"] - 1)).sum()
        exp = (st * bets["Expected profit"]).sum()
        header.append(f"The {nb} bet{'s' * (nb > 1)}: {st.sum():.2f} units staked to win {win:.2f} units; the model's "
                      f"expected profit {exp:+.2f} units.")
    quoted = [t for t in (local_time(x) for x in pred.get("quoted_at", pd.Series(dtype=object))) if t is not None]
    sources = sorted({_text(x) for x in pred.get("price_source", pd.Series(dtype=object)) if _text(x)})
    if quoted:
        q = max(quoted)
        age = (made - q).total_seconds() / 60
        header.append(f"Prices: {', '.join(s.capitalize() for s in sources) or 'the matchups CSV'}, quoted "
                      f"{when(q)} {zone} ({age:.0f} min before this sheet). Bet within "
                      f"{_f(meta.get('stale_minutes')):.0f} minutes of the quote, or reprice first (cell A6).")
    header.append(f"Model: {MODELS.get(meta.get('model'), meta.get('model') or '?')}; instance "
                  f"{meta.get('instance') or '?'}, seeds {meta.get('seeds') or '?'}.")
    status = _text(pred["patterns"].iloc[0]) if "patterns" in pred and len(pred) else ""
    if status == "on":
        sh = (pd.to_numeric(pred["q_a"], errors="coerce") -
              pd.to_numeric(pred.get("q_a_patterns_off"), errors="coerce")).abs() * 100
        biggest = float(sh.max()) if sh.notna().any() else 0.0
        eff = pred.get("patterns_effect", pd.Series(dtype=object))
        only = int((eff == "bet only with the patterns").sum())
        header.append(f"Pattern adjustment: ON. 'Model chance' includes it; 'Model chance without patterns' is the "
                      f"model before it. It moved chances by up to {biggest:.1f} pts here; {only} "
                      f"bet{'s' * (only != 1)} exist{'s' * (only == 1)} only because of it (Notes say which: keep those "
                      f"stakes small).")
    elif status == "off":
        header.append("Pattern adjustment: OFF. The model's chances are the version without it.")
    header.append("Stakes: " + (STAKE_RULES[stake_rule] + "." if stake_rule in STAKE_RULES else
                                "none shown (set STAKE_RULE in the notebook's setup cell)."))
    header.append("Expected profit is the model's own estimate at Pinnacle's odds, margin included. It is not a measured "
                  "return.")
    written = dt.datetime.now().astimezone()
    if (written - made).total_seconds() > 600:          # a sheet remade from a run's saved files, as of an earlier time
        header.append(f"Rebuilt {when(written, year=True)} {zone_label(written)} from the run's saved files, as of "
                      f"{when(made)} {zone}: the prices and blocks are those of that time.")

    run = [("Sheet made (as of)", f"{when(made, year=True)} {zone}"),
           ("Written", f"{when(written, year=True)} {zone_label(written)}"),
           ("Run", f"{meta.get('run_name', '')}  ({meta.get('run', '')})"),
           ("Model", MODELS.get(meta.get("model"), str(meta.get("model")))),
           ("Instance", str(meta.get("instance"))), ("State checkpoint", str(meta.get("checkpoint"))),
           ("Seeds", str(meta.get("seeds"))),
           ("Prices", f"{', '.join(sources)}; quoted {when(max(quoted)) if quoted else '?'} {zone}"),
           ("Reprice file", str(meta.get("reprice_file") or "none: the run's own prices")),
           ("Stale-price limit", f"{_f(meta.get('stale_minutes')):.0f} minutes"),
           ("Stake rule", STAKE_RULES.get(stake_rule or "", "none")),
           ("Pattern adjustment", {"on": "on (switched on 13 Oct 2026; docs/studies/unpriced_leads/)",
                                   "off": "off"}.get(status, "not recorded in this run")),
           ("Other venue", f"{venue} (fee {_f(meta.get('venue_fee')):.1%} of net winnings)" if venue else "none")]
    for i, nt in enumerate(notes, 1):
        run.append((f"Run note {i}", nt["plain"]))
        if nt.get("detail"):
            run.append(("", nt["detail"]))
    if meta.get("history_line") and not any(str(meta["history_line"]) in str(nt.get("detail")) for nt in notes):
        run.append(("History check", str(meta["history_line"])))
    return Sheet(made_at=made, zone=zone, title=title, header=header, notes=[nt["plain"] for nt in notes], table=table,
                 kinds=kinds, matches=matches, matches_kinds=mkinds, groups=groups, left_out=left,
                 left_out_summary=left_summary, run=run)


# ---------------------------------------------------------------------------------------------------------------------
# How to read: every column, plainly

def glossary(zone: str) -> list[tuple[str, str, str]]:
    S, M, N = "Bet sheet", "Matches", "Not predicted"
    return [
        (S, "(order)", "BET rows first, then PASS, then BLOCKED; each group by start time."),
        (S, f"Starts ({zone})", "The listed start in your time zone. Later matches on a court often start after it."),
        (S, "Match", "The two players as Pinnacle lists them (home first)."),
        (S, "Decision", "BET: the model sees positive expected profit at Pinnacle's own odds. PASS: no side does (the "
                        "model's edge is smaller than the bookmaker's margin). BLOCKED: a safety rule stops any bet; "
                        "Notes say which."),
        (S, "Player", "BET rows: the player to bet. PASS rows: the player the model rates above the market, not worth a "
                      "bet at Pinnacle's price."),
        (S, "Stake (units)", "The stake rule's size, in your units (see the header). Shown on BET rows only."),
        (S, "Odds", "Pinnacle's current decimal odds for that player, its margin included: the price you would get."),
        (S, "US odds", "The same price as American odds."),
        (S, "Lowest odds to take", "1 / model chance: below these odds the bet loses money on the model's numbers. "
                                   "Use it at any book or exchange (before its fees). It holds while Pinnacle's own "
                                   "price stays where it was: if Pinnacle moves, reprice (cell A6)."),
        (S, "Model chance", "The model's chance that this player wins. While the pattern adjustment is on (the "
                            "header says), this includes it, and the decision, edge and expected profit use it."),
        (S, "Model chance without patterns", "The same chance without the five-pattern adjustment (fast courts and "
                                             "big servers, rhythm, luck, the post-Slam slump, reputation; switched "
                                             "on 13 Oct 2026). When it makes the decision differ, Notes say so."),
        (S, "Market chance", "Pinnacle's price for this player with its margin removed (what the model is given)."),
        (S, "Edge (pts)", "Model chance minus market chance, in percentage points. Positive: the model rates the "
                          "player higher than the market does."),
        (S, "Expected profit", "Per unit staked at Pinnacle's odds: model chance x odds - 1. The model's own estimate, "
                               "not a measured return. It must clear the margin, so a small edge can still be a PASS."),
        (S, "Event", "Tour, event, round; surface and indoor or outdoor; best of 3 or 5."),
        (S, "Notes", "Why a match is a PASS or BLOCKED, and anything that weakens a bet (a thin edge, missing serve-"
                     f"speed readings, model runs that disagree). 'Thin' means expected profit under {THIN_EV:.0%}."),
        (M, "Player 1, Player 2", "Pinnacle's order (home first). Every '1' or '2' column is that player's."),
        (M, "Margin", "Pinnacle's overround: how much the two prices add up to over 100%. The bookmaker's cut."),
        (M, "Quoted, Price age (min)", "When Pinnacle's price was read, and how old it was when the sheet was made. A "
                                       "price older than the limit blocks the match until it is repriced."),
        (M, "Limit (max stake)", "The most Pinnacle accepts on this match (its maxRiskStake)."),
        (M, "Fair odds 1, 2", "1 / model chance: the model's own price for each player."),
        (M, "TabPFN alone 1", "Player 1's chance from TabPFN-3.5 before stage 2 (the average of its 16 runs)."),
        (M, "Stage 2 shift 1 (pts)", "How far stage 2 moved player 1's chance: its calibration (it pulls TabPFN's "
                                     "chances a little toward 50%) plus the serve-speed form term, plus the pattern "
                                     "adjustment while it is on."),
        (M, "Model chance 1 without patterns", "Player 1's chance without the five-pattern adjustment."),
        (M, "Pattern shift 1 (pts)", "How far the five-pattern adjustment moved player 1's chance (0 while it is off, "
                                     "or where a match has no pattern score)."),
        (M, "Pattern effect", "Whether the adjustment changes the decision: 'bet only with the patterns', 'the "
                              "patterns remove this bet', 'the patterns switch the side', or empty (no change)."),
        (M, "Serve-speed gap 1-2 (km/h)", "Player 1's recent average serve speed minus player 2's (28-day half-life), "
                                          "capped at about ±7.8 km/h. Faster favours that player. 0 for women (no feed) "
                                          "and when a player has no readings."),
        (M, "Speed readings (28 d)", "How many serve-speed readings each man has in the last 28 days."),
        (M, "Seed spread (± pts)", "How much the model's 16 runs disagree on this match (standard deviation). Above 3 "
                                   "pts is flagged."),
        (M, "Model's side", "The player the model rates above the market."),
        (M, "Expected profit 1, 2", "Each player's expected profit per unit at Pinnacle's odds."),
        (M, "Edge gone if the market reaches", "If Pinnacle's price for the model's side moved to this chance (these "
                                               "fair odds), the model would see no edge. 'Never' means the model prefers "
                                               "that player at every price from 5% to 95%."),
        (M, "Data check", "Complete or incomplete: the model reads the box scores of each player's earlier matches at "
                          "the event and in the last 14 days, and an incomplete match is blocked (Notes say whose "
                          "match lacks one)."),
        (M, "Notes", "The same notes as the bet sheet's."),
        (M, "Match id, Player ids, Run", "Identifiers in the project's data, for tracing a row."),
        (N, "Why not predicted", "Why a listed main-tour match has no prediction (Challengers, 125s, ITF, doubles and "
                                 "qualifying are left out before this list). 'Detail' is the linker's own text."),
        ("Run", "(all)", "Where the sheet came from: the run, the model instance, the prices, the rules, and the notes "
                         "that apply to every match."),
        ("records/", "(files)", "The machine record behind each sheet (every figure, sides A/B as Stage 0 fixes them) "
                                "and its provenance JSON; the daily engine's own archive of the run."),
    ]


# ---------------------------------------------------------------------------------------------------------------------
# the workbook

def write_workbook(sh: Sheet, path: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    def fill(c):
        return PatternFill("solid", start_color=c.lstrip("#"), end_color=c.lstrip("#"))

    thin = Side(style="thin", color="BFBFBF")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)

    def value(v):
        if v is None or (isinstance(v, float) and not np.isfinite(v)):
            return None
        if isinstance(v, (np.floating, np.integer)):
            v = v.item()
        return v

    def table(ws, df: pd.DataFrame, kinds: dict, top: int, *, colour_by: str | None = None) -> int:
        for j, c in enumerate(df.columns, 1):
            cell = ws.cell(row=top, column=j, value=c)
            cell.font, cell.fill = Font(bold=True, color=HEAD[1].lstrip("#")), fill(HEAD[0])
            cell.alignment, cell.border = Alignment(vertical="center", wrap_text=True), box
        for i, (_, row) in enumerate(df.iterrows(), top + 1):
            bg, fg = COLOURS.get(row.get(colour_by), (None, None)) if colour_by else (None, None)
            lines = 1
            for j, c in enumerate(df.columns, 1):
                cell = ws.cell(row=i, column=j, value=value(row[c]))
                fmt = FORMATS[kinds.get(c, "text")][0]
                if fmt:
                    cell.number_format = fmt
                cell.border = box
                cell.alignment = Alignment(vertical="top", wrap_text=c in WRAP)
                if c in WRAP and isinstance(row[c], str):        # Excel does not size wrapped rows by itself
                    lines = max(lines, math.ceil(len(row[c]) / (WRAP[c] * 1.15)))
                if bg:
                    cell.fill = fill(bg)
                    cell.font = Font(color=fg.lstrip("#"), bold=(c in ("Decision", "Player", "Bet on")
                                                                 and row.get(colour_by) == "BET"))
            if lines > 1:
                ws.row_dimensions[i].height = 15 * lines
        return top + len(df)

    def widths(ws, df: pd.DataFrame, kinds: dict, *, cap: dict | None = None, first: int = 1):
        cap = cap or {}
        for j, c in enumerate(df.columns, first):
            f = _fmt(kinds.get(c, "text")) or (lambda v: "" if value(v) is None else str(v))
            vals = [f(v) for v in df[c]]
            w = max([len(str(c)) * 0.9] + [len(v) for v in vals]) + 2
            ws.column_dimensions[get_column_letter(j)].width = WRAP[c] if c in WRAP else max(8, min(w, cap.get(c, 48)))

    wb = Workbook()
    # --- Bet sheet
    ws = wb.active
    ws.title = "Bet sheet"
    ws.cell(row=1, column=1, value=sh.title).font = Font(bold=True, size=14, color=HEAD[0].lstrip("#"))
    r = 2
    for line in sh.header:
        ws.cell(row=r, column=1, value=line).font = Font(size=11, bold=(r == 2))
        r += 1
    if sh.notes:
        ws.cell(row=r, column=1, value="Run notes (they apply to every match):").font = Font(bold=True)
        r += 1
        for nt in sh.notes:
            ws.cell(row=r, column=1, value=f"• {nt}")
            r += 1
    r += 1
    top = r
    end = table(ws, sh.table, sh.kinds, top, colour_by="Decision")
    widths(ws, sh.table, sh.kinds, cap={"Match": 46, "Event": 44})
    ws.row_dimensions[top].height = 30
    ws.freeze_panes = ws.cell(row=top + 1, column=1)
    if len(sh.table):
        ws.auto_filter.ref = f"A{top}:{get_column_letter(len(sh.table.columns))}{end}"
    if sh.left_out_summary:
        ws.cell(row=end + 2, column=1, value=sh.left_out_summary).font = Font(italic=True, color="595959")
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # --- Matches: group headers over column headers
    wm = wb.create_sheet("Matches")
    wm.cell(row=1, column=1, value="Every figure of each predicted match. Player 1 and 2 are in Pinnacle's order (home "
                                   "first); every '1' or '2' column is that player's.").font = Font(italic=True)
    col = 1
    for g, n in sh.groups:
        wm.merge_cells(start_row=2, start_column=col, end_row=2, end_column=col + n - 1)
        cell = wm.cell(row=2, column=col, value=g)
        cell.font, cell.fill = Font(bold=True, color=GROUP[1].lstrip("#")), fill(GROUP[0])
        cell.alignment = Alignment(horizontal="center")
        col += n
    end = table(wm, sh.matches, sh.matches_kinds, 3, colour_by="Decision")
    widths(wm, sh.matches, sh.matches_kinds, cap={"Match id": 50, "Edge gone if the market reaches": 44})
    wm.row_dimensions[3].height = 30
    wm.freeze_panes = wm.cell(row=4, column=4)              # the start and both players stay in view
    if len(sh.matches):
        wm.auto_filter.ref = f"A3:{get_column_letter(len(sh.matches.columns))}{end}"

    # --- Not predicted
    wn = wb.create_sheet("Not predicted")
    wn.cell(row=1, column=1, value=sh.left_out_summary or "Every listed main-tour match was predicted.").font = Font(
        italic=True)
    nk = {"Odds 1": "odds", "Odds 2": "odds"}
    if len(sh.left_out):
        end = table(wn, sh.left_out, nk, 3)
        widths(wn, sh.left_out, nk, cap={"Detail": 90, "Why not predicted": 70})
        wn.freeze_panes = wn.cell(row=4, column=1)
        wn.auto_filter.ref = f"A3:{get_column_letter(len(sh.left_out.columns))}{end}"

    # --- Run
    wr = wb.create_sheet("Run")
    for i, (k, v) in enumerate(sh.run + [("Files", "")] + [(k, str(v)) for k, v in sh.files.items()], 1):
        cell = wr.cell(row=i, column=1, value=k)
        cell.font, cell.alignment = Font(bold=True), Alignment(vertical="top")
        wr.cell(row=i, column=2, value=v).alignment = Alignment(wrap_text=True, vertical="top")
    wr.column_dimensions["A"].width, wr.column_dimensions["B"].width = 22, 140

    # --- How to read
    wg = wb.create_sheet("How to read")
    rows = glossary(sh.zone)
    for j, h in enumerate(("Sheet", "Column", "What it means"), 1):
        cell = wg.cell(row=1, column=j, value=h)
        cell.font, cell.fill = Font(bold=True, color=HEAD[1].lstrip("#")), fill(HEAD[0])
    for i, (s, c, m) in enumerate(rows, 2):
        wg.cell(row=i, column=1, value=s).alignment = Alignment(vertical="top")
        cell = wg.cell(row=i, column=2, value=c)
        cell.font, cell.alignment = Font(bold=True), Alignment(vertical="top")
        wg.cell(row=i, column=3, value=m).alignment = Alignment(wrap_text=True, vertical="top")
    wg.column_dimensions["A"].width, wg.column_dimensions["B"].width, wg.column_dimensions["C"].width = 14, 34, 120
    wg.freeze_panes = "A2"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# ---------------------------------------------------------------------------------------------------------------------
# the notebook's view

def _in_notebook() -> bool:
    try:
        from IPython import get_ipython
        ip = get_ipython()
    except ImportError:
        return False
    return ip is not None and type(ip).__name__ == "ZMQInteractiveShell"


def _fmt(kind: str):
    f = FORMATS[kind][1]
    return (lambda v: f.format(v) if isinstance(v, (int, float, np.floating)) and np.isfinite(v) else "") if f else None


def to_html(sh: Sheet) -> str:
    e = html.escape
    parts = [f"<div style='font-family:Segoe UI,Calibri,sans-serif'><h3 style='margin:4px 0'>{e(sh.title)}</h3>"]
    parts += [f"<div style='margin:2px 0;{'font-weight:600' if i == 0 else ''}'>{e(line)}</div>"
              for i, line in enumerate(sh.header)]
    if sh.notes:
        parts.append("<div style='margin-top:6px;font-weight:600'>Run notes (they apply to every match):</div><ul "
                     "style='margin:2px 0 6px 18px;padding:0'>" + "".join(f"<li>{e(n)}</li>" for n in sh.notes) + "</ul>")
    if len(sh.table):
        # the notebook's width: the event shortened, and each match's note listed under the table
        view = sh.table.drop(columns=["Notes"]).assign(Event=sh.table["Event"].str.split(" · ").str[0])
        view = view[[view.columns[0], "Event", *[c for c in view.columns[1:] if c != "Event"]]]
        sty = view.style.format({c: _fmt(k) for c, k in sh.kinds.items() if _fmt(k) and c in view}, na_rep="")
        sty = sty.apply(lambda row: [f"background-color:{COLOURS[row['Decision']][0]};color:"
                                     f"{COLOURS[row['Decision']][1]}" for _ in row], axis=1)
        sty = sty.hide(axis="index").set_table_styles([
            {"selector": "th", "props": f"background-color:{HEAD[0]};color:{HEAD[1]};padding:4px 6px;text-align:left"},
            {"selector": "td", "props": "padding:3px 6px;border-bottom:1px solid #d0d0d0;white-space:nowrap"},
            {"selector": "", "props": "border-collapse:collapse;font-size:12.5px;margin-top:6px"}])
        parts.append(sty.to_html())
        noted = sh.table[sh.table["Notes"].astype(str).str.len() > 0]
        if len(noted):
            items = "".join(
                f"<li><span style='background-color:{COLOURS[r['Decision']][0]};color:{COLOURS[r['Decision']][1]};"
                f"padding:0 4px;font-weight:600'>{e(r['Decision'])}</span> <b>{e(r['Match'])}</b>: {e(r['Notes'])}</li>"
                for _, r in noted.iterrows())
            parts.append(f"<ul style='margin:6px 0 2px 18px;padding:0'>{items}</ul>")
    else:
        parts.append("<div>No predicted match.</div>")
    if sh.left_out_summary:
        parts.append(f"<div style='margin-top:6px;color:#595959;font-style:italic'>{e(sh.left_out_summary)}</div>")
    if sh.files:
        parts.append("<div style='margin-top:6px'>" + "<br>".join(f"<b>{e(k)}:</b> {e(str(v))}"
                                                                  for k, v in sh.files.items()) + "</div>")
    return "".join(parts) + "</div>"


def to_text(sh: Sheet) -> str:
    out = [sh.title, *sh.header]
    if sh.notes:
        out += ["Run notes (they apply to every match):", *[f"  - {n}" for n in sh.notes]]
    if len(sh.table):
        t = sh.table.drop(columns=["Event", "Notes"])
        for c in t.columns:
            f = _fmt(sh.kinds.get(c, "text"))
            t[c] = t[c].map(f) if f else t[c].fillna("")
        with pd.option_context("display.width", 250, "display.max_columns", 40, "display.max_colwidth", 44):
            out += ["", t.to_string(index=False), ""]
        out += [f"  {r['Match']}: {r['Notes']}" for _, r in sh.table.iterrows() if r["Notes"]]
    if sh.left_out_summary:
        out += ["", sh.left_out_summary]
    out += [f"{k}: {v}" for k, v in sh.files.items()]
    return "\n".join(out)


def show(sh: Sheet) -> None:
    if _in_notebook():
        from IPython.display import HTML, display
        display(HTML(to_html(sh)))
    else:
        print(to_text(sh))
